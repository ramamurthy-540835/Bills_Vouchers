"""Bounded source-event normalization. Events never become invoices implicitly."""

import csv
import hashlib
import json
from io import StringIO, BytesIO
from decimal import Decimal
from itertools import islice
from typing import Any
from openpyxl import load_workbook

SOURCES = {"vendor", "customer", "driver"}
EVENTS = {"invoice", "credit_note", "debit_note", "trip", "payment", "settlement"}
FIELDS = {
    "source_record_id",
    "source_version",
    "event_type",
    "invoice_number",
    "invoice_date",
    "supplier_gstin",
    "customer_gstin",
    "trip_id",
    "booking_id",
    "vendor_id",
    "customer_id",
    "driver_id",
    "taxable_amount",
    "cgst",
    "sgst",
    "igst",
    "cess",
    "invoice_total",
    "payment_amount",
    "transaction_status",
}
MONEY = {"taxable_amount", "cgst", "sgst", "igst", "cess", "invoice_total", "payment_amount"}
MAX_ROWS = 10000
MAX_BYTES = 10 * 1024 * 1024


def normalize(source, record):
    if source not in SOURCES or not isinstance(record, dict):
        raise ValueError("Choose Vendor, Customer or Driver and supply an object.")
    if set(record) - FIELDS:
        raise ValueError("Unmapped fields; review the explicit mapping before import.")
    result: dict[str, Any] = {key: str(value).strip() for key, value in record.items() if value is not None}
    if not result.get("source_record_id") or len(result["source_record_id"]) > 200:
        raise ValueError("Source record ID is required (maximum 200 characters).")
    if result.get("event_type") not in EVENTS:
        raise ValueError("Choose a supported event type.")
    version = result.get("source_version", "1")
    if not version.isdigit() or not 1 <= int(version) <= 2**31 - 1:
        raise ValueError("Source version must be a positive integer.")
    result["source_version"] = int(version)
    for key in MONEY & result.keys():
        if isinstance(record[key], float):
            raise ValueError("Use decimal text for monetary fields.")
        value = Decimal(result[key])
        exponent = value.as_tuple().exponent
        if (
            not value.is_finite()
            or value < 0
            or not isinstance(exponent, int)
            or exponent < -2
            or value > Decimal("9999999999999999999999999999.99")
        ):
            raise ValueError("Monetary fields require finite nonnegative amounts with up to two decimals.")
        result[key] = str(value.quantize(Decimal(".01")))
    if any(len(str(value)) > 1000 for value in result.values()):
        raise ValueError("Field is too long.")
    result["source_system"] = source
    result["source_hash"] = hashlib.sha256(json.dumps(result, sort_keys=True).encode()).hexdigest()
    return result


def parse(content, filename, mapping=None):
    if len(content) > MAX_BYTES:
        raise ValueError("Maximum import size is 10 MB.")
    suffix = filename.lower().rsplit(".", 1)[-1]
    if suffix == "json":
        rows = json.loads(content)
    elif suffix == "csv":
        rows = list(islice(csv.DictReader(StringIO(content.decode("utf-8-sig"))), MAX_ROWS + 1))
    elif suffix == "xlsx":
        from zipfile import ZipFile

        with ZipFile(BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("Expanded workbook exceeds the safe size limit.")
        book = load_workbook(BytesIO(content), read_only=True, data_only=False)
        try:
            iterator = book.active.iter_rows(values_only=True)
            headings = next(iterator)
            rows = []
            for values in iterator:
                if len(rows) >= MAX_ROWS:
                    raise ValueError("Maximum 10000 rows per import.")
                rows.append(dict(zip(headings, (str(v) if isinstance(v, float) else v for v in values))))
        finally:
            book.close()
    else:
        raise ValueError("Use CSV, XLSX or JSON.")
    if not isinstance(rows, list) or len(rows) > MAX_ROWS:
        raise ValueError("Use an array with at most 10000 records.")
    if mapping:
        if len(set(mapping.values())) != len(mapping) or set(mapping.values()) - FIELDS:
            raise ValueError("Mapping destinations must be unique approved fields.")
        rows = [{target: row[source] for source, target in mapping.items() if source in row} for row in rows]
    return rows


def preview(source, rows):
    if source not in SOURCES:
        raise ValueError("Choose Vendor, Customer or Driver.")
    accepted = []
    errors = []
    seen: dict[tuple[str, int], str] = {}
    for index, row in enumerate(rows, 1):
        try:
            record = normalize(source, row)
            identity = (record["source_record_id"], record["source_version"])
            if identity in seen:
                status = "duplicate" if seen[identity] == record["source_hash"] else "conflict"
                errors.append({"row": index, "status": status, "message": "Repeated source identity in this batch."})
            else:
                seen[identity] = record["source_hash"]
                accepted.append(record)
        except (ValueError, ArithmeticError, TypeError):
            errors.append(
                {
                    "row": index,
                    "status": "quarantined",
                    "message": "Invalid or unmapped record. Check identity, event type, version and decimal amounts.",
                }
            )
    return {
        "accepted": accepted,
        "errors": errors,
        "financial_records_created": 0,
        "policy": "Staging only. Trips, payments and settlements never automatically create taxable invoices.",
    }
