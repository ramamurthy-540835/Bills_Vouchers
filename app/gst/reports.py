"""Read-only reports from canonical evidence; never a government filing adapter."""
import csv
import hashlib
import json
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO, StringIO
from xml.sax.saxutils import escape

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from app.services.gst.medallion import param, period_value
from app.services.gst.rules import HEADS
from . import returns
from .reconciliation import compare
from .report import calculate
from . import report_details, reconciliation_reviews

TITLES = {'purchases': 'Purchase register', 'sales': 'Sales register', 'invoices': 'Invoice listing',
          'tax': 'Tax-head summary', 'itc': 'ITC decisions', 'reconciliation': 'GSTR-2B reconciliation',
          'audit': 'Audit history', 'credit': 'Usable Input Tax Credit', 'liability': 'GST Liability on Sales',
          'cash': 'Estimated GST Cash Requirement', 'without_itc': 'Purchases Without ITC',
          'pending': 'Pending Bill Reviews', 'mismatches': 'Reconciliation Mismatches', 'unmatched': 'Unmatched Invoices'}
MAX_ROWS = 10000
NOTE = ('Canonical evidence only. Financial registers include confirmed invoices; invoice listing includes review states. '
        'Amounts are INR. Credit notes have negative signed values. Tax estimates are calculated separately for each month; '
        'they do not implement an opening/carry-forward credit ledger. This is a working report, not a GST filing export.')


def periods(period, scope):
    period = period_value(period)
    year, month = map(int, period.split('-'))
    if scope == 'month':
        return [period]
    if scope == 'quarter':
        first, count = (month-1)//3*3+1, 3
    elif scope == 'financial_year':
        year, first, count = year if month >= 4 else year-1, 4, 12
    else:
        raise HTTPException(422, 'Choose month, quarter or financial year.')
    if year < 2000 or year > 2099 or (scope == 'financial_year' and year == 2099):
        raise HTTPException(422, 'Choose a supported reporting year.')
    return [f'{year+(first-1+i)//12:04d}-{(first-1+i)%12+1:02d}' for i in range(count)]


def safe_cell(value):
    text = str(value if value is not None else '')
    text = ILLEGAL_CHARACTERS_RE.sub(lambda match: f'\\u{ord(match.group()):04x}', text)
    if len(text) > 32767:
        raise HTTPException(422, 'A report field exceeds the spreadsheet cell limit. Shorten the source field before exporting.')
    # Leading whitespace must not hide a spreadsheet formula.
    return "'" + text if text.lstrip().startswith(('=', '+', '-', '@')) or text.startswith(('\t', '\r', '\n')) else text


def audit_rows(store, period, cutoff):
    rows = store.query(f'''SELECT source,reference,actor_id,action,reason,created_at FROM (
      SELECT 'bill' AS source,COALESCE(invoice_id,JSON_VALUE(after_payload,'$.job_id'),'') AS reference,
        actor_id,action,COALESCE(JSON_VALUE(after_payload,'$.reason'),'') AS reason,created_at
      FROM {store.table('ops','gst_bill_audit')}
      WHERE client_id=@client_id AND created_at<=@cutoff AND (
        COALESCE(SUBSTR(JSON_VALUE(after_payload,'$.invoice_date'),1,7),JSON_VALUE(after_payload,'$.selected_period'),JSON_VALUE(after_payload,'$.period'))=@period OR
        COALESCE(SUBSTR(JSON_VALUE(before_payload,'$.invoice_date'),1,7),JSON_VALUE(before_payload,'$.selected_period'),JSON_VALUE(before_payload,'$.period'))=@period)
      UNION ALL
      SELECT 'reconciliation',case_id,actor_id,decision,reason,created_at
      FROM {store.table('ops','gst_reconciliation_review')} WHERE client_id=@client_id AND period=@period AND created_at<=@cutoff
      UNION ALL
      SELECT 'return',COALESCE(snapshot_id,''),actor_id,action,reason,created_at
      FROM {store.table('ops','gst_return_audit')} WHERE client_id=@client_id AND period=@period AND created_at<=@cutoff
    ) ORDER BY created_at,source,reference,action LIMIT {MAX_ROWS+1}''',
        [param('period', period), param('cutoff', cutoff, 'TIMESTAMP')])
    return [[period, *[str(row[key] or '') for key in ('source','reference','actor_id','action','reason','created_at')]] for row in rows]


def build(store, period, scope, kind, profile_for, filters=None):
    if kind not in TITLES:
        raise HTTPException(422, 'Choose a supported report.')
    selected = periods(period, scope)
    cutoff = datetime.now(timezone.utc)
    hashes = {month: returns.fingerprint(store, month) for month in selected}
    rows: list[list[str]] = []
    sources: list[list[str]] = []
    monthly_scenarios = []
    monthly_summary = []
    warnings = []
    detail_summary = []
    evidence_ids: list[str] = []
    columns = []
    def add(row):
        rows.append([str(value if value is not None else '') for value in row])
        if len(rows) > MAX_ROWS:
            raise HTTPException(422, 'This report exceeds 10,000 rows. Select a shorter period; no partial export was generated.')
    for month in selected:
        invoices = sorted(store.list(month), key=lambda inv: (str(inv.invoice_date or ''), inv.invoice_number, inv.invoice_id))
        if any(inv.client_id != store.client_id for inv in invoices):
            raise HTTPException(403, 'Report evidence does not belong to the selected business.')
        if any(inv.sample for inv in invoices):
            raise HTTPException(403, 'Synthetic invoices cannot be included in customer reports.')
        if len(invoices)+len(sources) > MAX_ROWS or sum(len(inv.line_items) for inv in invoices) > MAX_ROWS:
            raise HTTPException(422, 'This report is too large. Select a shorter period.')
        evidence_ids.extend(inv.invoice_id for inv in invoices if inv.extraction.raw_ref)
        sources.extend([[month, inv.invoice_id, str(inv.version), inv.status] for inv in invoices])
        if kind in {'credit','liability','cash','without_itc'}:
            calculation = calculate(invoices, profile_for(month), store.matching_rows(month))
            monthly_summary.append({'period':month, 'confirmed_count':calculation['confirmed_count'],
                'totals':{key:format(value,'.2f') for key,value in calculation['totals'].items()},
                'heads':{key:{head:format(value,'.2f') for head,value in values.items()} for key,values in calculation['heads'].items()}})
        if kind in report_details.KINDS:
            statements = store.matching_rows(month)
            imported = bool(statements) or bool(store.statement_history(month))
            reviews = reconciliation_reviews.latest(store,month) if kind in {'mismatches','unmatched'} else []
            columns, detail_rows, summary, messages = report_details.project(kind,invoices,profile_for(month),statements,imported,month,reviews)
            for row in detail_rows:
                add(row)
            warnings.extend(messages)
            detail_summary.append({'period':month,'values':summary})
            if kind == 'cash':
                monthly_scenarios.append({'period':month,**{key:format(value,'.2f') for key,value in calculation['scenarios'].items()}})
            continue
        if kind in {'purchases','sales','invoices','pending','liability'}:
            columns = ['Period','Invoice ID','Version','Status','Direction','Supplier GSTIN','Supplier','Recipient GSTIN','Recipient',
                       'Invoice number','Invoice date','Invoice type','Taxable value','IGST','CGST','SGST','Cess','Invoice total']
            for inv in invoices:
                if kind == 'pending' and inv.status not in {'needs_review','extracted'}:
                    continue
                if kind not in {'invoices','pending'} and (inv.status != 'confirmed' or inv.direction != ('purchase' if kind == 'purchases' else 'sale')):
                    continue
                sign = Decimal('-1') if inv.invoice_type == 'credit_note' else Decimal('1')
                add([month, inv.invoice_id, inv.version, inv.status, inv.direction, inv.supplier_gstin, inv.supplier_legal_name,
                     inv.recipient_gstin, inv.recipient_legal_name, inv.invoice_number, inv.invoice_date, inv.invoice_type,
                     *[format(sign*getattr(inv.totals, head), '.2f') for head in ('taxable_value', *HEADS, 'invoice_total')]])
        elif kind in {'tax','itc','credit','cash','without_itc'}:
            report = calculate(invoices, profile_for(month), store.matching_rows(month))
            monthly_scenarios.append({'period': month, **{key: format(value, '.2f') for key,value in report['scenarios'].items()}})
            if kind in {'tax','cash'}:
                columns = ['Period','Tax head','Eligible credit estimate','Output liability','Cash payable estimate']
                for head in HEADS:
                    add([month, head.upper(), *[format(report['heads'][name][head], '.2f') for name in ('credit','output','cash')]])
            else:
                columns = ['Period','Invoice ID','Invoice number','Supplier','Description','Matching','Eligibility','Rule','Signed GST']
                if kind == 'without_itc':
                    columns.append('Signed purchase value')
                purchase_lines = [(inv, line) for inv in invoices if inv.status == 'confirmed' and inv.direction == 'purchase' for line in inv.line_items]
                for position, row in enumerate(report['decisions']):
                    if kind == 'credit' and row['eligibility'] != 'eligible':
                        continue
                    if kind == 'without_itc' and row['eligibility'] == 'eligible':
                        continue
                    values = [month, *[row[key] for key in ('invoice_id','invoice_number','supplier_name','description','matching','eligibility','rule','tax')]]
                    if kind == 'without_itc':
                        inv, line = purchase_lines[position]
                        sign = Decimal('-1') if inv.invoice_type == 'credit_note' else Decimal('1')
                        values.append(format(sign * (line.taxable_value + sum((getattr(line, head) for head in HEADS), Decimal('0'))), '.2f'))
                    add(values)
        elif kind in {'reconciliation','mismatches','unmatched'}:
            columns = ['Period','Case ID','Invoice ID','Supplier GSTIN','Invoice number','Invoice date','Comparison','Candidate count',
                       'Taxable difference','IGST difference','CGST difference','SGST difference','Cess difference']
            for case in compare(invoices, store.matching_rows(month))['cases']:
                missing = case['status'] in {'missing_book','missing_statement'}
                if kind == 'unmatched' and not missing:
                    continue
                if kind == 'mismatches' and (missing or case['status'] == 'exact'):
                    continue
                add([month, case['id'], case['invoice_id'], case['supplier_gstin'], case['invoice_number'], case['invoice_date'],
                     case['status'], len(case['statement_candidates']),
                     *[(case['differences'] or {}).get(head) for head in ('taxable_value', *HEADS)]])
        else:
            columns = ['Period','Source','Reference','Actor ID','Action','Reason','UTC timestamp']
            for row in audit_rows(store, month, cutoff):
                add(row)
    if hashes != {month: returns.fingerprint(store, month) for month in selected}:
        raise HTTPException(409, 'Source evidence changed while generating this report. Please retry.')
    result = {'kind': kind, 'title': TITLES[kind], 'client_id': store.client_id, 'periods': selected,
              'generated_at': cutoff.isoformat(), 'columns': columns, 'rows': rows, 'source_invoices': sources,
              'evidence_ids':evidence_ids,'warnings':list(dict.fromkeys(warnings)), 'detail_summary':detail_summary,
              'business_name':profile_for(period).get('legal_name',''), 'gstin':profile_for(period).get('gstin',''),
              'source_hashes': hashes, 'monthly_cash_scenarios': monthly_scenarios, 'monthly_summary': monthly_summary, 'note': NOTE, 'row_count': len(rows)}
    apply_filters(result, filters or {})
    result['report_hash'] = hashlib.sha256(json.dumps(result, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    return result


FILTER_COLUMNS = {'party':['Supplier','Recipient'], 'invoice_type':['Invoice type'], 'eligibility':['Eligibility'],
                  'matching':['Matching','Comparison'], 'source':['Source']}
NUMERIC_COLUMNS = {'Taxable value','IGST','CGST','SGST','Cess','Invoice total','Signed GST','Signed purchase value'}


def apply_filters(report, filters):
    """Filter authorized bounded evidence before pagination and export, with no SQL interpolation."""
    from datetime import date
    columns, original = report['columns'], report['rows']
    for key in ('date_from','date_to'):
        if filters.get(key):
            try:date.fromisoformat(filters[key])
            except ValueError:raise HTTPException(422,'Use an ISO invoice date (YYYY-MM-DD).') from None
    if filters.get('date_from') and filters.get('date_to') and filters['date_from']>filters['date_to']:
        raise HTTPException(422,'Start date must not follow end date.')
    if any(len(str(value))>200 for value in filters.values()):
        raise HTTPException(422,'Report filters must be at most 200 characters.')
    report['filter_options'] = {key:sorted({row[columns.index(column)] for column in names if column in columns for row in original if row[columns.index(column)]}) for key,names in FILTER_COLUMNS.items()}
    def keep(row):
        if filters.get('q') and filters['q'].casefold() not in ' '.join(row).casefold():return False
        for key,names in FILTER_COLUMNS.items():
            if filters.get(key) and not any(column in columns and filters[key]==row[columns.index(column)] for column in names):return False
        for key in ('date_from','date_to'):
            if filters.get(key):
                value = row[columns.index('Invoice date')] if 'Invoice date' in columns else ''
                if not value or (key=='date_from' and value<filters[key]) or (key=='date_to' and value>filters[key]):return False
        return True
    rows = [row for row in original if keep(row)]
    sort = filters.get('sort','')
    if sort:
        column = sort.removeprefix('-')
        if column not in columns:raise HTTPException(422,'Choose a report column to sort.')
        position = columns.index(column)
        rows.sort(key=lambda row:Decimal(row[position] or '0') if column in NUMERIC_COLUMNS else row[position].casefold(),reverse=sort.startswith('-'))
    report['rows'], report['row_count'], report['filters'] = rows,len(rows),filters
    if filters and 'Invoice ID' in columns:
        selected_ids = {row[columns.index('Invoice ID')] for row in rows}
        report['source_invoices'] = [row for row in report['source_invoices'] if row[1] in selected_ids]
    report['summary'] = {column:format(sum((Decimal(row[index] or '0') for row in rows),Decimal('0')),'.2f') for index,column in enumerate(columns) if column in NUMERIC_COLUMNS}
    report['note'] += ' Row totals and exports respect filters. Monthly tax and cash context describes the full period and is not recalculated from a filtered subset.'


def render(report, format):
    if format not in {'csv','xlsx','pdf'} or (format == 'pdf' and report['kind'] != 'tax'):
        raise HTTPException(422, 'Choose CSV or Excel. PDF is available for the tax-head summary.')
    metadata = [['Report', report['title']], ['Business',report.get('business_name','')], ['GSTIN',report.get('gstin','')], ['Client ID', report['client_id']], ['Periods', ', '.join(report['periods'])],
                ['Generated UTC', report['generated_at']], ['Report SHA256', report['report_hash']], ['Scope', report['note']], ['Filters', json.dumps(report.get('filters',{}),sort_keys=True)],
                ['Dependencies','; '.join(report.get('warnings',[]))]]
    context = [['Month','Tax head','Confirmed invoices','Usable credit','Output GST','Estimated cash']]
    for summary in report.get('monthly_summary', []):
        for head in HEADS:
            context.append([summary['period'],head.upper(),str(summary['confirmed_count']),
                *[summary['heads'][key][head] if summary['confirmed_count'] else 'Not calculated' for key in ('credit','output','cash')]])
    if format == 'csv':
        output = StringIO()
        writer = csv.writer(output)
        writer.writerows([[safe_cell(value) for value in row] for row in [*metadata, [], report['columns'], *report['rows']]])
        writer.writerow([])
        writer.writerow(['Source period','Source fingerprint'])
        writer.writerows(report['source_hashes'].items())
        if len(context)>1:
            writer.writerow([])
            writer.writerows(context)
        if report['monthly_cash_scenarios']:
            writer.writerow([])
            writer.writerow(['Month','Recorded credit','Matched credit','Eligibility checked','No credit'])
            writer.writerows([[row[key] for key in ('period','recorded','matched','eligible','noCredit')] for row in report['monthly_cash_scenarios']])
        return output.getvalue().encode('utf-8-sig'), 'text/csv'
    if format == 'xlsx':
        book = Workbook()
        sheet = book.active
        sheet.title = 'Report'
        for row in [report['columns'], *report['rows']]:
            sheet.append([safe_cell(value) for value in row])
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        meta = book.create_sheet('Provenance')
        for row in [*metadata, [], ['Period','Source fingerprint'], *report['source_hashes'].items()]:
            meta.append([safe_cell(value) for value in row])
        evidence = book.create_sheet('Source invoices')
        for row in [['Period','Invoice ID','Version','Status'], *report['source_invoices']]:
            evidence.append([safe_cell(value) for value in row])
        if report['monthly_cash_scenarios']:
            scenarios = book.create_sheet('Monthly cash scenarios')
            scenarios.append(['Period','Recorded credit','Matched credit','Eligibility checked','No credit'])
            for row in report['monthly_cash_scenarios']:
                scenarios.append([row[key] for key in ('period','recorded','matched','eligible','noCredit')])
        if len(context)>1:
            sheet_context = book.create_sheet('Financial context')
            for row in context:
                sheet_context.append([safe_cell(value) for value in row])
        output_bytes = BytesIO()
        book.save(output_bytes)
        return output_bytes.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    return pdf_summary(report, metadata), 'application/pdf'


def pdf_summary(report, metadata):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    buffer = BytesIO()
    styles = getSampleStyleSheet()
    # PDF contains English summary labels/amounts/identifiers; full Unicode registers use CSV/XLSX.
    def paragraph(value, style='BodyText'):
        text = str(value).encode('ascii', 'backslashreplace').decode()
        return Paragraph(escape(text), styles[style])
    story = [paragraph('GST Easy - tax-head summary', 'Title')]
    for key,value in metadata:
        story.append(paragraph(f'{key}: {value}'))
    story.append(Spacer(1, 14))
    table = Table([[paragraph(value) for value in row] for row in [report['columns'], *report['rows']]],
                  colWidths=[85, 65, 170, 160, 170], repeatRows=1)
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e5eef9')),
                              ('GRID',(0,0),(-1,-1),0.3,colors.grey),('VALIGN',(0,0),(-1,-1),'TOP')]))
    story.append(table)
    story.append(Spacer(1, 14))
    story.append(paragraph('Monthly cash estimates - four independent scenarios', 'Heading2'))
    for row in report['monthly_cash_scenarios']:
        story.append(paragraph(f"{row['period']}: recorded {row['recorded']}; matched {row['matched']}; eligibility checked {row['eligible']}; no credit {row['noCredit']} INR"))
    story.append(paragraph('Source fingerprints', 'Heading2'))
    for month, digest in report['source_hashes'].items():
        story.append(paragraph(f'{month}: {digest}'))
    def footer(canvas, doc):
        canvas.setFont('Helvetica', 8)
        canvas.drawString(36, 20, 'Working report - not a GST filing export')
        canvas.drawRightString(805, 20, f'Page {doc.page}')
    SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=36, rightMargin=36,
                      topMargin=30, bottomMargin=35).build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()
