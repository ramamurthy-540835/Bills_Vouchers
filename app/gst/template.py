import csv
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO, StringIO

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from .schemas.invoice import GstInvoice, LineItem, Totals, Extraction
from .validate_invoice import RATES
from app.services.gst.validation import VALID_STATE_CODES

HEADERS = ['Supplier GSTIN', 'Invoice Number', 'Invoice Date (DD-MM-YYYY)', 'HSN/SAC',
           'Description', 'Taxable Value', 'GST Rate %', 'CGST', 'SGST', 'IGST', 'Cess',
           'Invoice Total', 'Place of Supply (State Code)', 'Reverse Charge (Y/N)']


def template_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Bills'
    sheet.append(HEADERS)
    sheet.freeze_panes = 'A2'
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 26
    instructions = workbook.create_sheet('Instructions')
    for row in [
        ['One row per line item. Repeat invoice identity and full Invoice Total for each line.'],
        ['Use dates DD-MM-YYYY. Enter GST amounts in rupees. Do not include currency symbols.'],
        ['The recipient is your selected business. Supplier name and addresses can be completed in review.'],
        ['Do not change headers. Identical invoice identities are grouped into one bill.'],
        ['A missing business GSTIN permits upload but prevents confirmation.'],
        ['Allowed rates', ', '.join(str(r) for r in sorted(RATES))],
        ['State codes', ', '.join(sorted(VALID_STATE_CODES))],
        ['Common states', '07 Delhi; 27 Maharashtra; 29 Karnataka; 33 Tamil Nadu; 36 Telangana; 37 Andhra Pradesh'],
    ]:
        instructions.append(row)
    instructions.column_dimensions['A'].width = 110
    instructions.column_dimensions['B'].width = 110
    for col, values in [('G', sorted(RATES)), ('M', sorted(VALID_STATE_CODES)), ('N', ['Y', 'N'])]:
        validation = DataValidation(type='list', formula1='"' + ','.join(map(str, values)) + '"')
        validation.errorTitle = 'Choose a listed value'
        validation.error = 'Select a value from the dropdown.'
        validation.showErrorMessage = True
        sheet.add_data_validation(validation)
        validation.add(f'{col}2:{col}10001')
    dates = DataValidation(type='date', operator='between', formula1='DATE(2017,7,1)', formula2='DATE(2100,12,31)')
    dates.showErrorMessage = True
    sheet.add_data_validation(dates)
    dates.add('C2:C10001')
    for row in range(2, 1002):
        sheet.cell(row, 3).number_format = 'DD-MM-YYYY'
        for col in (1, 2, 4, 13):
            sheet.cell(row, col).number_format = '@'
    out = BytesIO()
    workbook.save(out)
    return out.getvalue()


def parse_template(content, extension, client_id, doc_id, profile, raw_ref):
    if extension == '.csv':
        rows = list(csv.reader(StringIO(content.decode('utf-8-sig'))))
    else:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            rows = list(workbook['Bills'].values)
        finally:
            workbook.close()
    if not rows or list(rows[0]) != HEADERS:
        raise ValueError('Use the provided Excel template or its exact CSV headers.')
    if len(rows) > 10001:
        raise ValueError('Upload at most 10,000 rows per file.')
    groups = {}
    for values in rows[1:]:
        if not any(v is not None and str(v).strip() for v in values):
            continue
        row = dict(zip(HEADERS, values, strict=True))
        day = row[HEADERS[2]]
        if isinstance(day, datetime):
            day = day.date()
        elif not isinstance(day, date):
            day = datetime.strptime(str(day), '%d-%m-%Y').date()
        supplier = str(row['Supplier GSTIN'] or '').strip().upper()
        number = str(row['Invoice Number'] or '').strip()
        key = (supplier, number, day)
        dec = lambda name: Decimal(str(row[name] if row[name] is not None else '0'))
        line = LineItem(description=str(row['Description'] or ''), hsn_sac=str(row['HSN/SAC'] or ''),
                        taxable_value=dec('Taxable Value'), gst_rate=dec('GST Rate %'),
                        cgst=dec('CGST'), sgst=dec('SGST'), igst=dec('IGST'), cess=dec('Cess'))
        if key not in groups:
            groups[key] = GstInvoice(client_id=client_id, source_doc_id=doc_id,
                supplier_gstin=supplier, recipient_gstin=profile.get('gstin') or '',
                recipient_legal_name=profile.get('legal_name') or '',
                recipient_address=profile.get('principal_address') or '',
                invoice_number=number, invoice_date=day,
                place_of_supply_state_code=str(row['Place of Supply (State Code)'] or '').zfill(2),
                reverse_charge=str(row['Reverse Charge (Y/N)'] or '').upper() == 'Y',
                extraction=Extraction(method='csv' if extension == '.csv' else 'xlsx_template', raw_ref=raw_ref),
                totals=Totals(invoice_total=dec('Invoice Total')))
        inv = groups[key]
        if inv.totals.invoice_total != dec('Invoice Total'):
            raise ValueError('Repeat the same invoice total on every line of a bill.')
        inv.line_items.append(line)
    for inv in groups.values():
        for field in ('taxable_value', 'cgst', 'sgst', 'igst', 'cess'):
            setattr(inv.totals, field, sum((getattr(line, field) for line in inv.line_items), Decimal('0')))
    return list(groups.values())
