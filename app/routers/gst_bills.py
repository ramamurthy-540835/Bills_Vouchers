import csv
import logging
import mimetypes
import hashlib
from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any
from uuid import uuid4, uuid5, NAMESPACE_URL

from fastapi import APIRouter, Depends, HTTPException, Request, UploadFile, File, Query
from fastapi.responses import Response, JSONResponse
from openpyxl import Workbook
from pydantic import ValidationError
from google.api_core.exceptions import InvalidArgument, PermissionDenied, FailedPrecondition

from app.db import get_db
from app.config import get_settings
from app.services.demo_identity import isolated_demo_workspace
from app.repository import FinanceRepository
from app.routes import active_client, current_user
from app.services.gst.account import can
from app.services.gst.medallion import Medallion, clean, period_value
from app.gst.schemas.invoice import GstInvoice, LineItem, Totals
from app.gst.store import BillStore
from app.gst.naming import bill_filename
from app.gst.template import template_bytes
from app.gst.validate_invoice import validate_invoice
from app.gst.file_checks import validate_file
from app.gst.report import calculate
from app.gst.reconciliation import compare, FIELDS
from app.gst import reconciliation_reviews, reports, dashboard as dashboard_report
from app.gst.matching import parse_gstr2b
from app.gst.jobs import enqueue, run_job, public_job
from app.services.tasks import verify_cloud_tasks_request

router = APIRouter(prefix='/api/gst/bills', tags=['Bills'])
logger = logging.getLogger(__name__)


def scope(request, repo, user, action='bills:read'):
    if not get_settings().gst_bills_enabled:
        raise HTTPException(404, 'Not found.')
    client = active_client(request, repo, user)
    if not can(user, action, client.id, repo):
        raise HTTPException(403, 'You do not have permission for this action.')
    readonly_demo_report = isolated_demo_workspace() and request.url.path.endswith('/reports') and request.query_params.get('format', 'json') == 'json'
    if getattr(client, 'is_test', False) and not readonly_demo_report and request.url.path.rsplit('/', 1)[-1] in {'export', 'working-papers', 'reports'}:
        raise HTTPException(403, {'code': 'sample_data_blocked', 'message': 'Test workspace figures cannot be exported or filed.'})
    return client, BillStore(repo, client.id)


def profile_for(repo, client_id, period):
    return Medallion(repo, client_id, period).profile() or {}


def selected_period(request, client, store, explicit=''):
    chosen = explicit or request.session.get('period')
    if not chosen and getattr(client, 'is_test', False):
        chosen = store.latest_period()
    return period_value(chosen or date.today().strftime('%Y-%m'))


def safe_cell(value):
    value = str(value or '')
    return "'" + value if value.startswith(('=', '+', '-', '@', '\t', '\r')) else value


def check_version(request, invoice):
    expected = request.headers.get('if-match')
    if expected and expected.strip('"') != str(invoice.version):
        raise HTTPException(409, 'This bill changed in another window. Reopen it before saving.')


@router.get('/template.xlsx')
def template(request: Request, repo=Depends(get_db), user=Depends(current_user)):
    scope(request, repo, user)
    return Response(template_bytes(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition': "attachment; filename*=UTF-8''Bills%20%26%20Vouchers%20%E2%80%94%20GST%20Purchase%20Bills.xlsx"})


@router.post('/upload')
def upload(request: Request, files: list[UploadFile] = File(...), period: str = '', auto_review: bool = False, repo=Depends(get_db), user=Depends(current_user)):
    client, store = scope(request, repo, user, 'bills:write')
    period = selected_period(request, client, store, period)
    store.assert_writable(period)
    profile = profile_for(repo, client.id, period)
    if not 1 <= len(files) <= 20:
        raise HTTPException(422, 'Choose between 1 and 20 bills.')
    content_files: list[tuple[str, bytes]] = []
    allowed = {'.pdf', '.jpg', '.jpeg', '.png', '.xlsx', '.csv'}
    for file in files:
        filename = Path((file.filename or 'bill').replace('\\', '/')).name
        if Path(filename).suffix.lower() not in allowed:
            raise HTTPException(422, 'Choose PDF, JPG, PNG, Excel or CSV files.')
        content = file.file.read(get_settings().max_upload_mb * 1024 * 1024 + 1)
        if not content or len(content) > get_settings().max_upload_mb * 1024 * 1024:
            raise HTTPException(413, f'Each bill must be between 1 byte and {get_settings().max_upload_mb} MB.')
        if sum(len(c) for _, c in content_files) + len(content) > 25 * 1024 * 1024:
            raise HTTPException(413, 'Upload smaller batches, up to 25 MB per request.')
        try:
            validate_file(content, filename)
        except (ValueError, OSError):
            raise HTTPException(422, 'Check the file format and size. Use a valid PDF, photo or the provided spreadsheet template.') from None
        content_files.append((filename, content))
    job_id = str(uuid4())
    job: dict[str, Any] = {'job_id': job_id, 'status': 'queued', 'recovery_version': 1, 'actor_id': str(user.id), 'period': period, 'files': [], 'auto_review_enabled': auto_review}
    for filename, content in content_files:
        doc_id = str(uuid5(NAMESPACE_URL, f'{client.id}/{hashlib.sha256(content).hexdigest()}'))
        raw_ref = store.upload_object(content, doc_id, filename, 'bronze', mimetypes.guess_type(filename)[0] or 'application/octet-stream')
        job['files'].append({'filename': filename, 'doc_id': doc_id, 'raw_ref': raw_ref, 'status': 'reading', 'invoice_ids': []})
    store.job(job_id, job)
    if get_settings().gst_bills_queue:
        try:
            enqueue(str(client.id), job_id)
        except (InvalidArgument, PermissionDenied, FailedPrecondition):
            logger.exception('Bill job could not be queued')
            job['status'] = 'needs_attention'
            for entry in job['files']:
                entry.update(status='needs_attention', reason='Your original is saved, but reading could not start. Please retry this upload.')
            store.job(job_id, job)
            raise HTTPException(503, 'Your original bills were saved, but reading could not start. Please retry shortly.') from None
        except Exception:
            logger.exception('Bill enqueue outcome unknown',extra={'job_id':job_id})
            raise HTTPException(503, 'Your original bills are saved. Dispatch could not be confirmed; check upload history before uploading again.') from None
    else:
        job=run_job(store, profile, str(user.id), job)
    return public_job(job)


@router.post('/internal/process')
async def worker(request: Request, repo=Depends(get_db)):
    settings = get_settings()
    if not settings.gst_bills_enabled or not verify_cloud_tasks_request(request, settings.gst_bills_worker_audience or settings.gst_bills_worker_url, settings.gst_bills_worker_account):
        raise HTTPException(403, 'Access denied.')
    payload = await request.json()
    store = BillStore(repo, str(payload['client_id']))
    job = store.job(str(payload['job_id']))
    actor = FinanceRepository(repo).user_by_id(job['actor_id'])
    if actor:
        from app.services.gst.medallion import param
        from app.services.access import workspace_role
        role = repo.one(f"SELECT access_role AS role FROM `{repo.table('client_memberships')}` WHERE client_id=@client_id AND user_id=@uid AND is_active=TRUE ORDER BY created_at DESC LIMIT 1", [param('client_id',store.client_id),param('uid',str(actor.id))])
        actor.role = workspace_role(role.role if role else None)
    if not actor or not can(actor, 'bills:write', store.client_id, repo):
        raise HTTPException(403, 'The uploader no longer has permission to save these bills.')
    from starlette.concurrency import run_in_threadpool
    await run_in_threadpool(run_job, store, profile_for(repo,store.client_id,job['period']), str(actor.id), job)
    return {'completed': True}


@router.get('/upload/{job_id}')
def progress(job_id: str, request: Request, repo=Depends(get_db), user=Depends(current_user)):
    _, store = scope(request, repo, user)
    return public_job(store.job(job_id))


@router.get('/history')
def upload_history(request: Request, period: str = '', repo=Depends(get_db), user=Depends(current_user)):
    client, store = scope(request, repo, user)
    return {'jobs':[public_job(job) for job in store.job_history(selected_period(request,client,store,period))]}


@router.post('/upload/{job_id}/retry')
def retry_upload(job_id: str, request: Request, repo=Depends(get_db), user=Depends(current_user)):
    client,store=scope(request,repo,user,'bills:write')
    job=store.retry_job(job_id,str(user.id))
    if get_settings().gst_bills_queue:
        try:
            enqueue(str(client.id),job_id,uuid4().hex)
        except Exception:
            # A timeout can mean Cloud Tasks accepted the request. Leave the job queued
            # rather than racing a possibly active worker with an unconditional write.
            logger.exception('Bill retry enqueue outcome unknown',extra={'job_id':job_id})
            raise HTTPException(503, 'Retry dispatch could not be confirmed. Your original is retained; check upload history before retrying.') from None
    else:
        job=run_job(store,profile_for(repo,client.id,job['period']),str(user.id),job)
    return public_job(job)


@router.post('/sum-lines')
async def sum_lines(request: Request, repo=Depends(get_db), user=Depends(current_user)):
    scope(request,repo,user,'bills:write')
    try:
        body=await request.json()
        values=body.get('line_items') if isinstance(body,dict) else None
        if not isinstance(values,list) or not 1 <= len(values) <= 2000:
            raise ValueError('Invalid line count')
        lines=[LineItem.model_validate(value) for value in values]
        totals={head:sum((getattr(line,head) for line in lines),Decimal('0')) for head in ('taxable_value','cgst','sgst','igst','cess')}
        result=Totals(**totals,invoice_total=sum(totals.values(),Decimal('0')))
    except (ValueError,ValidationError,InvalidOperation):
        raise HTTPException(422, 'Enter valid decimal amounts for between 1 and 2000 line items.') from None
    return result.model_dump(mode='json')


@router.get('/deleted')
def deleted_bills(request: Request, period: str = '', repo=Depends(get_db), user=Depends(current_user)):
    client, store = scope(request,repo,user,'bills:write')
    return {'invoices':[{**inv.model_dump(mode='json'),'filename':bill_filename(inv)} for inv in store.deleted(selected_period(request,client,store,period))]}


@router.post('/{invoice_id}/restore')
def restore_bill(invoice_id: str, request: Request, repo=Depends(get_db), user=Depends(current_user)):
    client, store = scope(request,repo,user,'bills:write')
    previous = store.deleted_invoice(invoice_id)
    check_version(request,previous)
    period = previous.invoice_date.strftime('%Y-%m') if previous.invoice_date else previous.selected_period
    invoice = validate_invoice(previous,profile_for(repo,client.id,period),store.list())
    invoice.status = 'needs_review'
    invoice.auto_review.status = 'not_started'
    invoice.auto_review.job_id = ''
    invoice.auto_review.reasons = []
    invoice.auto_review.checks = {}
    invoice.version = previous.version+1
    store.save(invoice,user.id,'restore',selected_period=period)
    return invoice.model_dump(mode='json')


@router.get('')
def list_bills(request: Request, period: str = '', status: str = '', repo=Depends(get_db), user=Depends(current_user)):
    client, store = scope(request, repo, user)
    period = selected_period(request, client, store, period)
    return {'invoices': [{**i.model_dump(mode='json'), 'filename': bill_filename(i)} for i in store.list(period, status)],
            'client_id': str(client.id), 'can_write': can(user, 'bills:write', client.id, repo)}


@router.get('/dashboard')
def dashboard(request: Request, period: str = '', scope_name: str = 'month', repo=Depends(get_db), user=Depends(current_user)):
    client, store = scope(request, repo, user)
    period = selected_period(request, client, store, period)
    calculation = dashboard_report.build(store, period, scope_name, lambda month: profile_for(repo, client.id, month))
    return JSONResponse(clean({'client_id': str(client.id), 'period': period, 'sample': False,
        'test_workspace': bool(getattr(client, 'is_test', False)),
        'user': {'id': str(user.id), 'full_name': user.full_name, 'email': user.email, 'role': user.role},
        'clients': [{'id': c.id, 'name': c.name} for c in FinanceRepository(repo).clients(user.id)],
        'can_write': can(user, 'bills:write', client.id, repo),
        **calculation}), headers={'Cache-Control': 'private, no-store'})


@router.get('/corrections')
def correction_queue(request: Request, period: str = '', state: str = 'active',
                     offset: int = Query(0, ge=0, le=100000),
                     repo=Depends(get_db), user=Depends(current_user)):
    client, store = scope(request, repo, user)
    if state not in {'active','rejected','correction_requested','corrected','resubmitted','approved'}:
        raise HTTPException(422, 'Choose a supported correction state.')
    period = selected_period(request, client, store, period)
    result = store.correction_queue(period, state, offset)
    return JSONResponse(clean({**result, 'period': period, 'state': state,
        'next_offset': offset+50 if result['has_more'] else None}),
        headers={'Cache-Control': 'private, no-store'})


@router.post('/gstr2b')
async def import_gstr2b(request: Request, period: str, repo=Depends(get_db), user=Depends(current_user)):
    client, store = scope(request, repo, user, 'bills:write')
    period = period_value(period)
    content = await request.body()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(413, 'Choose a statement smaller than 10 MB.')
    try:
        records = parse_gstr2b(content, profile_for(repo, client.id, period), period)
    except (ValueError, TypeError, KeyError, InvalidOperation):
        raise HTTPException(422, 'Check that this is the GSTR-2B JSON for your GSTIN and selected period, without repeated invoices.') from None
    store.assert_writable(period)
    digest=hashlib.sha256(content).hexdigest()
    ref=store.upload_object(content,str(uuid5(NAMESPACE_URL,f'{client.id}/gstr2b/{digest}')),f'GSTR2B-{period}-{digest[:12]}.json','bronze','application/json')
    source={'sha256':digest,'period':period,'raw_ref':ref,'parser_version':'b2b-2','record_count':len(records)}
    store.save_matching(period, records, user.id,source)
    return {'imported': len(records)}


@router.get('/reconciliation')
def reconciliation(request: Request, period: str, repo=Depends(get_db), user=Depends(current_user)):
    client, store = scope(request, repo, user)
    period = period_value(period)
    return clean({'period': period, 'can_review': user.role in {'admin','tax_admin'} and not getattr(client,'is_test',False),
                  **reconciliation_reviews.view(store, period)})


@router.post('/reconciliation/review')
async def record_reconciliation_review(request: Request, period: str, repo=Depends(get_db), user=Depends(current_user)):
    client, store = scope(request, repo, user, 'bills:write')
    if getattr(client, 'is_test', False):
        raise HTTPException(403, 'Synthetic workspaces cannot record financial review decisions.')
    return reconciliation_reviews.record(store, period_value(period), user, await request.json())


@router.get('/reconciliation/history')
def reconciliation_history(request: Request, period: str, case_id: str, before_version: int | None = Query(None,ge=1,le=2**63-1), repo=Depends(get_db), user=Depends(current_user)):
    _, store = scope(request, repo, user)
    return clean(reconciliation_reviews.history(store, period_value(period), case_id, before_version))


@router.get('/gstr2b/history')
def statement_history(request: Request,period: str,repo=Depends(get_db),user=Depends(current_user)):
    _,store=scope(request,repo,user)
    return {'statements':[{key:value for key,value in row.items() if key!='raw_ref'} for row in store.statement_history(period_value(period))]}


@router.get('/gstr2b/{digest}/original')
def statement_original(digest: str,request: Request,period: str,repo=Depends(get_db),user=Depends(current_user)):
    _,store=scope(request,repo,user)
    period=period_value(period)
    row=next((row for row in store.statement_history(period) if row['sha256']==digest),None)
    if row is None:
        raise HTTPException(404,'Statement not found.')
    return Response(store.download_object(row['raw_ref']),media_type='application/json',
        headers={'Content-Disposition':f'attachment; filename="GSTR2B-{period}.json"','Cache-Control':'private, no-store'})


@router.get('/working-papers')
def working_papers(request: Request, period: str, repo=Depends(get_db), user=Depends(current_user)):
    client, store = scope(request, repo, user)
    period = period_value(period)
    invoices = store.list(period)
    if any(i.sample for i in invoices):
        raise HTTPException(403, {'code': 'sample_data_blocked', 'message': 'Illustrations cannot be exported.'})
    statement = store.matching_rows(period)
    report = calculate(invoices, profile_for(repo, client.id, period), statement)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'GST Easy summary'
    sheet.append(['Period', period])
    labels = {'eligible_credit': 'Usable Input Tax Credit', 'output_tax': 'GST Liability on Sales',
              'cash_required': 'Estimated GST Cash Requirement', 'non_gst': 'Purchases Without ITC',
              'recorded': 'As recorded in your bills', 'matched': 'After GSTR-2B Matching',
              'eligible': 'After ITC Eligibility Checks', 'noCredit': 'If No ITC Can Be Used'}
    for key, value in report['totals'].items():
        sheet.append([labels[key], str(value)])
    for key, value in report['scenarios'].items():
        sheet.append([labels[key], str(value)])
    decisions = workbook.create_sheet('Credit decisions')
    decisions.append(['Invoice', 'Supplier', 'Description', 'Matching', 'Eligibility', 'Rule', 'GST'])
    for row in report['decisions']:
        decisions.append([safe_cell(row[k]) for k in ('invoice_number', 'supplier_name', 'description', 'matching', 'eligibility', 'rule', 'tax')])
    reconciliation_sheet = workbook.create_sheet('Reconciliation evidence')
    reconciliation_sheet.append(['Comparison only; no manual credit approval. Differences = books minus statement.'])
    reconciliation_sheet.append(['Invoice ID', 'Supplier GSTIN', 'Invoice', 'Book date', 'Status',
                                 'Candidate count', *['Difference ' + head for head in FIELDS]])
    comparison = compare(invoices, statement)
    for case in comparison['cases']:
        reconciliation_sheet.append([safe_cell(case[key]) for key in
                                     ('invoice_id', 'supplier_gstin', 'invoice_number', 'invoice_date', 'status')]
                                    + [len(case['statement_candidates'])]
                                    + [str(case['differences'][head]) if case['differences'] and case['differences'][head] is not None else '' for head in FIELDS])
    sources = workbook.create_sheet('Statement candidates')
    sources.append(['Case ID', 'Statement row', 'Invoice', 'Date', 'Source SHA256', 'ITC available', *FIELDS])
    for case in comparison['cases']:
        for candidate in case['statement_candidates']:
            sources.append([safe_cell(case['id']), candidate['position'],
                            *[safe_cell(candidate[key]) for key in ('invoice_number', 'invoice_date', 'source_sha256')],
                            str(candidate['itc_available']),
                            *[str(candidate['amounts'][head]) if candidate['amounts'][head] is not None else '' for head in FIELDS]])
    reviews = workbook.create_sheet('Reconciliation reviews')
    reviews.append(['Workflow review only; not an ITC override. Revalidate against current evidence before use.'])
    reviews.append(['Case ID', 'Version', 'Decision', 'Reason', 'Reviewer ID', 'Reviewed at', 'Evidence hash'])
    for review in reconciliation_reviews.latest(store, period):
        reviews.append([safe_cell(review[key]) for key in ('case_id','version','decision','reason','actor_id','created_at','source_hash')])
    output = BytesIO()
    workbook.save(output)
    return Response(output.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': f'attachment; filename="GST-Easy-CA-working-papers-{period}.xlsx"'})


@router.get('/export')
def export(request: Request, period: str, format: str = 'xlsx', repo=Depends(get_db), user=Depends(current_user)):
    _, store = scope(request, repo, user)
    invoices = store.list(period_value(period))
    if any(i.sample for i in invoices):
        raise HTTPException(403, {'code': 'sample_data_blocked', 'message': 'Illustrations cannot be exported.'})
    headers = ['GSTIN of supplier', 'Trade/Legal name', 'Invoice number', 'Invoice type', 'Invoice Date', 'Invoice Value',
               'Place of supply', 'Supply Attract Reverse Charge', 'Taxable Value', 'Integrated Tax', 'Central Tax', 'State/UT Tax', 'Cess', 'ITC Availability']
    rows = [headers]
    for inv in invoices:
        if inv.direction != 'purchase' or inv.status != 'confirmed':
            continue
        rows.append([safe_cell(inv.supplier_gstin), safe_cell(inv.supplier_legal_name), safe_cell(inv.invoice_number),
                     inv.invoice_type, str(inv.invoice_date), str(inv.totals.invoice_total), inv.place_of_supply_state_code,
                     'Y' if inv.reverse_charge else 'N', *[str(getattr(inv.totals, f)) for f in ('taxable_value', 'igst', 'cgst', 'sgst', 'cess')], inv.itc.eligibility])
    if format == 'csv':
        stream = StringIO()
        csv.writer(stream).writerows(rows)
        data, mime = stream.getvalue().encode('utf-8-sig'), 'text/csv'
    elif format == 'xlsx':
        workbook = Workbook()
        workbook.active.title = 'Purchase register'
        for row in rows:
            workbook.active.append(row)
        buffer = BytesIO()
        workbook.save(buffer)
        data, mime = buffer.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    else:
        raise HTTPException(422, 'Choose Excel or CSV.')
    return Response(data, media_type=mime, headers={'Content-Disposition': f'attachment; filename="GST-Easy-purchase-register-{period}.{format}"'})


@router.get('/reports')
def reports_view(request: Request, period: str, kind: str = 'purchases', scope_name: str = 'month',
                 format: str = 'json', offset: int = Query(0, ge=0, le=10000),
                 repo=Depends(get_db), user=Depends(current_user)):
    client, store = scope(request, repo, user)
    if getattr(client, 'is_test', False) and not (isolated_demo_workspace() and format == 'json'):
        raise HTTPException(403, 'Synthetic workspaces cannot generate customer reports.')
    if format not in {'json','csv','xlsx','pdf'} or (format == 'pdf' and kind != 'tax'):
        raise HTTPException(422, 'Choose CSV or Excel; PDF is available for the tax-head summary.')
    filters = {key:request.query_params[key] for key in ('q','party','invoice_type','eligibility','matching','source','date_from','date_to','sort') if request.query_params.get(key)}
    report = reports.build(store, period, scope_name, kind, lambda month: profile_for(repo, client.id, month), filters)
    private = {'Cache-Control': 'private, no-store', 'X-Content-Type-Options': 'nosniff'}
    if format == 'json':
        from fastapi.responses import JSONResponse
        return JSONResponse({**report, 'can_export': not getattr(client, 'is_test', False), 'rows': report['rows'][offset:offset+100],
                             'source_invoices': report['source_invoices'][offset:offset+100],
                             'source_invoice_count': len(report['source_invoices']),
                             'offset': offset, 'has_more': offset+100 < report['row_count']}, headers=private)
    content, mime = reports.render(report, format)
    filename = f"GST-Easy-{kind}-{report['periods'][0]}-to-{report['periods'][-1]}.{format}"
    return Response(content, media_type=mime, headers={**private,'Content-Disposition':f'attachment; filename="{filename}"'})


@router.get('/{invoice_id}/original')
def original(invoice_id: str, request: Request, repo=Depends(get_db), user=Depends(current_user)):
    _, store = scope(request, repo, user)
    invoice = store.get(invoice_id)
    ref = invoice.extraction.raw_ref
    mime = mimetypes.guess_type(ref)[0] or 'application/octet-stream'
    return Response(store.download_object(ref), media_type=mime,
                    headers={'Content-Disposition': f'inline; filename="{bill_filename(invoice)}"', 'Cache-Control': 'private, no-store'})


@router.get('/{invoice_id}')
def get_bill(invoice_id: str, request: Request, repo=Depends(get_db), user=Depends(current_user)):
    _, store = scope(request, repo, user)
    invoice = store.get(invoice_id)
    locked = ''
    try:
        store.assert_writable(invoice.invoice_date.strftime('%Y-%m') if invoice.invoice_date else invoice.selected_period)
    except HTTPException as exc:
        if exc.status_code != 409:
            raise
        locked = 'This filing period is locked; Auto Review is unavailable.'
    result = invoice.model_dump(mode='json')
    if invoice.auto_review.status == 'running' and invoice.auto_review.job_id:
        job = store.job(invoice.auto_review.job_id)
        if job['status'] == 'needs_attention' or public_job(job)['retryable']:
            result['auto_review']['status'] = 'failed'
            result['auto_review']['reasons'] = ['Auto Review stopped. Run Auto Review to retry the saved invoice, or turn OFF.']
        elif job['status'] == 'completed' and job.get('result') == 'cancelled':
            result['auto_review']['status'] = 'cancelled'
            result['auto_review']['reasons'] = ['The invoice changed or its period was locked. Refresh before trying again.']
    return {**result, 'filename': bill_filename(invoice), 'auto_review_locked': locked}


@router.delete('/{invoice_id}')
def delete_bill(invoice_id: str, request: Request, repo=Depends(get_db), user=Depends(current_user)):
    _, store = scope(request, repo, user, 'bills:write')
    invoice = store.get(invoice_id)
    check_version(request, invoice)
    store.delete(invoice, user.id)
    return {'deleted': True, 'invoice_id': invoice_id}


@router.patch('/{invoice_id}')
async def edit(invoice_id: str, request: Request, repo=Depends(get_db), user=Depends(current_user)):
    client, store = scope(request, repo, user, 'bills:write')
    previous = store.get(invoice_id)
    check_version(request, previous)
    body = await request.json()
    owned = {'client_id', 'invoice_id', 'source_doc_id', 'selected_period', 'extraction', 'itc', 'status', 'review_reasons', 'rejection_reason', 'sample', 'version', 'supply_type',
             'rejection_code','correction_state','correction_reason','correction_actor','correction_at','auto_review'}
    if owned.intersection(body):
        raise HTTPException(422, 'Only bill details may be edited.')
    if 'reverse_charge_paid_reference' in body and user.role not in {'admin', 'tax_admin'}:
        raise HTTPException(403, 'A tax administrator must record reverse-charge payment evidence.')
    try:
        invoice = GstInvoice.model_validate({**previous.model_dump(), **body, 'version': previous.version+1})
    except ValidationError:
        raise HTTPException(422, 'Check the bill fields and enter amounts as decimal text.') from None
    period = invoice.invoice_date.strftime('%Y-%m') if invoice.invoice_date else previous.selected_period or date.today().strftime('%Y-%m')
    invoice = validate_invoice(invoice, profile_for(repo, client.id, period), store.list())
    if invoice.auto_review.enabled:
        invoice.auto_review.status = 'not_started'
        invoice.auto_review.job_id = ''
        invoice.auto_review.checks = {}
        invoice.auto_review.reasons = []
        invoice.auto_review.rule_ids = []
        invoice.auto_review.completed_at = ''
    if 'duplicate_invoice' in invoice.review_reasons:
        raise HTTPException(409, 'This supplier GSTIN and invoice number already exist for the financial year.')
    if previous.status=='rejected':
        if previous.correction_state not in {'correction_requested','corrected'}:
            raise HTTPException(409,'Request correction before editing a rejected bill.')
        invoice.status='rejected'
        invoice.correction_state='corrected'
    store.save(invoice, user.id, 'edit', previous, period)
    return invoice.model_dump(mode='json')


@router.post('/{invoice_id}/confirm')
def confirm(invoice_id: str, request: Request, repo=Depends(get_db), user=Depends(current_user)):
    client, store = scope(request, repo, user, 'bills:write')
    previous = store.get(invoice_id)
    check_version(request, previous)
    if previous.sample:
        raise HTTPException(403, {'code': 'sample_data_blocked', 'message': 'Illustrations cannot be confirmed.'})
    period = (previous.invoice_date or date.today()).strftime('%Y-%m')
    invoice = validate_invoice(previous, profile_for(repo, client.id, period), store.list(), human_review=True)
    if invoice.review_reasons or previous.status == 'rejected':
        raise HTTPException(409, {'code': 'bill_needs_review', 'message': 'Resolve the highlighted bill details before confirming.'})
    invoice.status = 'confirmed'
    if previous.correction_state=='resubmitted':
        invoice.correction_state='approved'


    if invoice.auto_review.status == 'running':
        invoice.auto_review.status = 'cancelled'
        invoice.auto_review.reasons = ['An authorized user confirmed this bill manually.']
    invoice.version += 1
    store.save(invoice, user.id, 'confirm', previous, period)
    return invoice.model_dump(mode='json')


@router.post('/{invoice_id}/reject')
async def reject(invoice_id: str, request: Request, repo=Depends(get_db), user=Depends(current_user)):
    _, store = scope(request, repo, user, 'bills:write')
    previous = store.get(invoice_id)
    check_version(request, previous)
    try:
        body = await request.json()
    except ValueError:
        raise HTTPException(422, 'Enter a rejection reason.') from None
    reason = body.get('reason') if isinstance(body, dict) else None
    if not isinstance(reason, str) or not 3 <= len(reason.strip()) <= 1000:
        raise HTTPException(422, 'Enter a rejection reason between 3 and 1000 characters.')
    from app.gst.correction import REASONS
    code=body.get('code','other')
    if not isinstance(code,str) or code not in REASONS:
        raise HTTPException(422,'Choose a supported rejection reason.')
    invoice = previous.model_copy(deep=True)
    invoice.status = 'rejected'
    if invoice.auto_review.status == 'running':
        invoice.auto_review.status = 'cancelled'
        invoice.auto_review.reasons = ['An authorized user rejected this bill.']
    invoice.rejection_reason = reason.strip()
    invoice.rejection_code=code
    invoice.correction_state='rejected'
    invoice.version += 1
    store.save(invoice, user.id, 'reject', previous, (previous.invoice_date or date.today()).strftime('%Y-%m'))
    return invoice.model_dump(mode='json')


@router.get('/{invoice_id}/gst-comparison')
def gst_comparison(invoice_id: str, request: Request, repo=Depends(get_db), user=Depends(current_user)):
    from app.gst.comparison import compare_invoice
    from app.services.gst.medallion import param
    import json
    _,store=scope(request,repo,user)
    invoice=store.get(invoice_id)
    rows=store.query(f'''SELECT after_payload FROM {store.table('ops','gst_bill_audit')}
        WHERE client_id=@client_id AND invoice_id=@id
          AND (before_payload IS NULL OR JSON_TYPE(before_payload)='null')
        ORDER BY created_at LIMIT 1''',[param('id',invoice_id)])
    payload=rows[0]['after_payload'] if rows else None
    source=GstInvoice.model_validate(json.loads(payload) if isinstance(payload,str) else payload) if payload else None
    return compare_invoice(invoice,source)


@router.get('/{invoice_id}/review-history')
def review_history(invoice_id: str, request: Request, repo=Depends(get_db), user=Depends(current_user)):
    from app.services.gst.medallion import param
    _,store=scope(request,repo,user)
    store.get(invoice_id)
    rows=store.query(f'''SELECT action,actor_id,created_at,
        JSON_VALUE(after_payload,'$.version') AS version,
        JSON_VALUE(after_payload,'$.rejection_reason') AS rejection_reason,
        JSON_VALUE(after_payload,'$.correction_reason') AS correction_reason
        FROM {store.table('ops','gst_bill_audit')}
        WHERE client_id=@client_id AND invoice_id=@id ORDER BY created_at DESC LIMIT 100''',[param('id',invoice_id)])
    return clean({'events':[dict(row) for row in rows]})


@router.post('/{invoice_id}/correction/{action}')
async def correction_action(invoice_id: str, action: str, request: Request, repo=Depends(get_db), user=Depends(current_user)):
    from app.gst.correction import transition
    client,store=scope(request,repo,user,'bills:write')
    previous=store.get(invoice_id)
    if not request.headers.get('if-match'):
        raise HTTPException(428,'Refresh the bill and submit its current version.')
    check_version(request,previous)
    body=await request.json()
    invoice=transition(previous,action,user.id,body.get('reason') if isinstance(body,dict) else None)
    period=previous.invoice_date.strftime('%Y-%m') if previous.invoice_date else previous.selected_period
    if action=='resubmit':
        validated=validate_invoice(invoice,profile_for(repo,client.id,period),store.list(),human_review=True)
        if 'duplicate_invoice' in validated.review_reasons:
            raise HTTPException(409,'Resolve the duplicate before resubmitting.')
        invoice.review_reasons=validated.review_reasons
    store.save(invoice,user.id,action,previous,period)
    return invoice.model_dump(mode='json')


def auto_review_target(invoice_id, request, repo, user):
    client, store = scope(request, repo, user, 'bills:write')
    invoice = store.get(invoice_id)
    if not request.headers.get('if-match'):
        raise HTTPException(428, 'Reopen the invoice before changing Auto Review.')
    check_version(request, invoice)
    if invoice.status == 'rejected' or invoice.sample:
        raise HTTPException(409, 'Auto Review is unavailable for rejected or illustration invoices.')
    period = invoice.invoice_date.strftime('%Y-%m') if invoice.invoice_date else invoice.selected_period
    store.assert_writable(period)
    return client, store, invoice, period


@router.patch('/{invoice_id}/auto-review')
async def auto_review_policy(invoice_id: str, request: Request, repo=Depends(get_db), user=Depends(current_user)):
    from app.gst.auto_review import now
    _, store, previous, period = auto_review_target(invoice_id, request, repo, user)
    body = await request.json()
    if not isinstance(body, dict) or set(body) != {'enabled'} or type(body['enabled']) is not bool:
        raise HTTPException(422, 'Choose Auto Review ON or OFF.')
    if previous.status == 'confirmed' and body['enabled']:
        raise HTTPException(409, 'This invoice is already confirmed; use the existing correction workflow.')
    invoice = previous.model_copy(deep=True)
    invoice.auto_review.enabled = body['enabled']
    invoice.auto_review.requested_by = str(user.id)
    invoice.auto_review.requested_at = now()
    if previous.status != 'confirmed':
        invoice.auto_review.status = 'not_started' if body['enabled'] else 'cancelled'
        invoice.auto_review.reasons = []
        invoice.auto_review.checks = {}
        invoice.auto_review.rule_ids = []
        invoice.auto_review.completed_at = ''
        invoice.auto_review.job_id = ''
    invoice.version += 1
    store.save(invoice, user.id, 'auto_review_policy', previous, period)
    return invoice.model_dump(mode='json')


@router.post('/{invoice_id}/auto-review/run')
def start_auto_review(invoice_id: str, request: Request, repo=Depends(get_db), user=Depends(current_user)):
    client, store, previous, period = auto_review_target(invoice_id, request, repo, user)
    if not previous.auto_review.enabled or previous.status == 'confirmed':
        raise HTTPException(409, 'Enable Auto Review on an unconfirmed invoice first.')
    if previous.auto_review.status == 'running':
        prior_job = store.job(previous.auto_review.job_id)
        if prior_job['status'] not in {'needs_attention', 'completed'} and not public_job(prior_job)['retryable']:
            raise HTTPException(409, 'Auto Review is already queued or running. Refresh its status.')
    invoice = previous.model_copy(deep=True)
    invoice.version += 1
    job_id = str(uuid4())
    invoice.auto_review.status = 'running'
    invoice.auto_review.job_id = job_id
    job = {'job_id': job_id, 'kind': 'auto_review', 'status': 'queued', 'recovery_version': 1,
           'actor_id': str(user.id), 'period': period, 'files': [],
           'invoice_id': invoice_id, 'invoice_version': invoice.version}
    store.job(job_id, job)
    store.save(invoice, user.id, 'auto_review_requested', previous, period)
    if get_settings().gst_bills_queue:
        try:
            enqueue(str(client.id), job_id)
        except Exception:
            logger.error('Auto Review dispatch uncertain', extra={'job_id': job_id})
            raise HTTPException(503, 'Auto Review dispatch is unconfirmed. Refresh; do not submit again. Saved work is retained.') from None
    else:
        run_job(store, profile_for(repo, client.id, period), str(user.id), job)
    return store.get(invoice_id).model_dump(mode='json')
