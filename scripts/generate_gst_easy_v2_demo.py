"""Reproducible synthetic documents only. Never calls an application/cloud API."""
import hashlib
import json
from copy import deepcopy
from decimal import Decimal as D
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont, ImageFilter
from openpyxl import Workbook
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

ROOT = Path('demo bills')
NOTICE = 'SYNTHETIC TEST INVOICE - NOT VALID FOR TAX FILING'
PERIOD = '2026-10'
HEADS = ('taxable_value', 'cgst', 'sgst', 'igst', 'cess', 'invoice_total')
BASE36 = '0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ'


def gstin(state, index):
    base = state + f'ZXZCZ{index:04d}Q' + '1Z'
    total, factor = 0, 2
    for char in reversed(base):
        value = BASE36.index(char) * factor
        total += value // 36 + value % 36
        factor = 3 - factor
    return base + BASE36[(36 - total % 36) % 36]


BUSINESSES = [dict(key=k, legal_name='Nature Labs Demo '+title+' Pvt Ltd', trade_name='Nature Labs '+title,
    state_code=state, city=city, pin=pin, contact_email=k+'@example.test', contact_phone='+91900000010'+str(i+1),
    website_url='https://'+k+'.example.test', gstin=gstin(state,i+1), pan=gstin(state,i+1)[2:12],
    constitution='Company', registration_type='Regular', filing_frequency='monthly', business_nature=activity,
    principal_address=f'Synthetic Demo Campus, Unit {i+1}, {city}, {pin} (fictional address)',
    is_test=True, gstin_verified=False, provenance='SYNTHETIC_TEST_ONLY')
    for i,(k,title,state,city,pin,activity) in enumerate([
        ('manufacturing','Manufacturing','33','Chennai','600032','Manufacturing of industrial components'),
        ('trading','Trading','33','Coimbatore','641018','Wholesale and Retail Trading'),
        ('services','Services','29','Bengaluru','560001','Professional Services')])]
DESCRIPTIONS = {
 'manufacturing':['Raw material steel','Copper wire','Electrical components','Industrial machine spares','Packaging materials','Factory equipment maintenance','Industrial control panel','Electrical distribution board','Fabricated metal component','Industrial sensor assembly','Machinery spare parts','Finished electrical equipment'],
 'trading':['Wholesale electrical procurement','Industrial materials','Tools procurement','Equipment procurement','Retail electrical goods','Wholesale materials','Industrial tools credit note','Interstate equipment'],
 'services':['IT supplier services','Technical maintenance inward','Installation subcontractor','Software support inward','IT consulting','Technical maintenance outward','Equipment installation','Software support debit note']}


def money(value):
    return str(D(value).quantize(D('.01')))


def words(n):
    small='zero one two three four five six seven eight nine ten eleven twelve thirteen fourteen fifteen sixteen seventeen eighteen nineteen'.split()
    tens='zero ten twenty thirty forty fifty sixty seventy eighty ninety'.split()
    if n<20:return small[n]
    if n<100:return tens[n//10]+(' '+small[n%10] if n%10 else '')
    for value,label in [(10000000,'crore'),(100000,'lakh'),(1000,'thousand'),(100,'hundred')]:
        if n>=value:return words(n//value)+' '+label+(' '+words(n%value) if n%value else '')
    raise ValueError(n)


def font(size,bold=False):
    for path in [Path('C:/Windows/Fonts/arialbd.ttf' if bold else 'C:/Windows/Fonts/arial.ttf'),Path('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf' if bold else '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')]:
        if path.exists():return ImageFont.truetype(str(path),size)
    return ImageFont.load_default(size=size)


def invoice(business,index):
    count=len(DESCRIPTIONS[business['key']]);purchase=index<count//2
    other_state='29' if business['state_code']=='33' else '33'
    counter_state=other_state if index%3==0 else business['state_code']
    counter=dict(legal_name=f'Synthetic {"Supplier" if purchase else "Customer"} {index+1:02}',gstin=gstin(counter_state,100+index),state_code=counter_state,principal_address=f'Fictional Industrial Road {index+1}, Test City')
    supplier,recipient=(counter,business) if purchase else (business,counter)
    inter=supplier['state_code']!=recipient['state_code']; lines=[]
    for j in range(2):
        qty=D(j+2);unit=D(500+(index+1)*137+j*47);taxable=qty*unit;rate=D(18 if j==0 else 12)
        tax=(taxable*rate/100).quantize(D('.01'))
        lines.append(dict(description=DESCRIPTIONS[business['key']][index]+(' - accessory/support' if j else ''),hsn_sac='998313' if business['key']=='services' else '853710',quantity=str(qty),unit='NOS',unit_price=money(unit),taxable_value=money(taxable),gst_rate=str(rate),cgst=money(0 if inter else tax/2),sgst=money(0 if inter else tax/2),igst=money(tax if inter else 0),cess='0.00'))
    totals={h:money(sum((D(line[h]) for line in lines),D(0))) for h in HEADS[:-1]}
    totals['invoice_total']=money(sum((D(v) for v in totals.values()),D(0)))
    note='credit_note' if business['key']=='trading' and index==6 else 'debit_note' if business['key']=='services' and index==7 else 'tax_invoice'
    return dict(business=business['key'],invoice_number=f'NL-{business["key"][:3].upper()}-{index+1:03}',invoice_date=f'2026-10-{index+1:02}',invoice_type=note,direction='purchase' if purchase else 'sale',supplier=supplier,recipient=recipient,place_of_supply_state_code=recipient['state_code'],line_items=lines,totals=totals,provenance='SYNTHETIC_TEST_ONLY',rate_assumption='18% and 12% illustrative test rates; HSN/SAC classification not legally certified',itc_assumption='Pending review until a user establishes eligibility; no automatic ITC approval',reconciliation=('match' if index==0 else 'amount_mismatch' if index==1 else 'tax_head_mismatch' if index==2 else 'books_only') if purchase else 'not_applicable')


def render(inv):
    img=Image.new('RGB',(1240,1754),'white');d=ImageDraw.Draw(img)
    d.rectangle((0,0,1240,130),fill='#123e30');d.text((45,30),'NATURE LABS / SYNTHETIC GST '+inv['invoice_type'].replace('_',' ').upper(),font=font(30,True),fill='white')
    d.text((45,92),NOTICE,font=font(23,True),fill='white')
    y=175
    lines=[f'Invoice: {inv["invoice_number"]}     Date: {inv["invoice_date"]}',f'Direction: {inv["direction"].upper()} | Place of supply: {inv["place_of_supply_state_code"]}']
    for role in ['supplier','recipient']:
        person=inv[role];lines += [role.title()+': '+person['legal_name'],'GSTIN-like: '+person['gstin'],person['principal_address']]
    for line in lines:d.text((45,y),line,font=font(25),fill='#163c30');y+=44
    y+=30
    d.text((45,y),'DESCRIPTION / HSN-SAC / QTY / UNIT PRICE / TAXABLE / RATE',font=font(22,True),fill='#123e30');y+=55
    for line in inv['line_items']:
        d.text((45,y),line['description'],font=font(23,True),fill='black');y+=36
        d.text((45,y),f'{line["hsn_sac"]} | {line["quantity"]} {line["unit"]} | INR {line["unit_price"]} | INR {line["taxable_value"]} | {line["gst_rate"]}%',font=font(22),fill='black');y+=35
        d.text((45,y),'  '.join(h.upper()+': '+line[h] for h in ['cgst','sgst','igst','cess']),font=font(22),fill='#444444');y+=58
    for h in HEADS:
        d.text((650,y),h.replace('_',' ').title()+': INR '+inv['totals'][h],font=font(25,True),fill='#123e30');y+=44
    total=D(inv['totals']['invoice_total']);amount_words='Rupees '+words(int(total))+' and '+words(int(total%1*100))+' paise only'
    import textwrap
    for line in textwrap.wrap(amount_words,75):d.text((45,y+30),line,font=font(23),fill='black');y+=35
    for line in ['Payment reference: DEMO-NO-PAYMENT / No real bank account','Reference: Synthetic order '+inv['invoice_number'],'Test rates only; not tax advice. Credit notes reduce signed totals.','Fictional authorized signatory: Nature Labs Demo','SYNTHETIC TEST INVOICE - NOT VALID FOR TAX FILING']:
        y+=45;d.text((45,y),line,font=font(22),fill='#555555')
    return img


def pdf(image,path,pages=1):
    c=canvas.Canvas(str(path),pagesize=(620,877),invariant=1)
    for _ in range(pages):c.drawImage(ImageReader(image),0,0,620,877);c.showPage()
    c.save()


def main():
    ROOT.mkdir(exist_ok=True);(ROOT/'edge_cases').mkdir(exist_ok=True);(ROOT/'gstr2b').mkdir(exist_ok=True);(ROOT/'logos').mkdir(exist_ok=True)
    manifest={'provenance':'SYNTHETIC_TEST_ONLY','period':PERIOD,'provisioning':'profiles only; never import these files automatically','businesses':BUSINESSES,'invoices':[],'edge_cases':[]}
    wb=Workbook();sheet=wb.active;sheet.title='Invoices'
    sheet.append(['Business','Invoice identity','Type','Direction','Supplier','Recipient','Date',*HEADS,'Expected validation','Duplicate','Reconciliation','ITC','Dashboard contribution','Report contribution'])
    totals_sheet=wb.create_sheet('Expected signed totals');totals_sheet.append(['Business','Direction',*HEADS])
    for b in BUSINESSES:
        folder=ROOT/b['key'];folder.mkdir(exist_ok=True)
        logo=Image.new('RGB',(320,320),{'manufacturing':'#187653','trading':'#1f5385','services':'#63418a'}[b['key']]);d=ImageDraw.Draw(logo);d.rounded_rectangle((25,25,295,295),radius=35,outline='white',width=5);d.text((72,78),'NL',font=font(100,True),fill='white');d.text((40,215),b['key'].upper(),font=font(25,True),fill='white');logo.save(ROOT/'logos'/f'{b["key"]}.png')
        docs=[invoice(b,i) for i in range(len(DESCRIPTIONS[b['key']]))]
        for inv in docs:
            image=render(inv);stem=inv['supplier']['gstin']+'_'+inv['invoice_number'];paths=[]
            for ext in ['pdf','png','jpg']:
                path=folder/(stem+'.'+ext)
                if ext=='pdf':pdf(image,path)
                else:image.save(path)
                paths.append(str(path.relative_to(ROOT)).replace('\\','/'))
            inv['files']=paths;manifest['invoices'].append(inv)
            sign=D(-1) if inv['invoice_type']=='credit_note' else D(1)
            sheet.append([b['key'],inv['invoice_number'],inv['invoice_type'],inv['direction'],inv['supplier']['legal_name'],inv['recipient']['legal_name'],inv['invoice_date'],*[str(D(inv['totals'][h])*sign) for h in HEADS],'Review required; structurally valid','Same business/FY/identity duplicate',inv['reconciliation'],inv['itc_assumption'],'Only after manual confirmation; signed amounts above','Relevant register and GST detail after confirmation'])
        for direction in ['purchase','sale']:
            totals_sheet.append([b['key'],direction,*[money(sum((D(i['totals'][h])*(-1 if i['invoice_type']=='credit_note' else 1) for i in docs if i['direction']==direction),D(0))) for h in HEADS]])
        statement={'gstin':b['gstin'],'rtnprd':'102026','docdata':{'b2b':[]}}
        for index,inv in enumerate([i for i in docs if i['direction']=='purchase'][:3]):
            t=inv['totals'];row={'inum':inv['invoice_number'],'dt':'-'.join(reversed(inv['invoice_date'].split('-'))),'itcavl':'Y',**{out:t[src] for out,src in [('txval','taxable_value'),('cgst','cgst'),('sgst','sgst'),('igst','igst'),('cess','cess')]}}
            if index==1:row['cgst']=money(D(row['cgst'])+10)
            if index==2:row['igst']=money(D(row['cgst'])+D(row['sgst']));row['cgst']=row['sgst']='0.00'
            statement['docdata']['b2b'].append({'ctin':inv['supplier']['gstin'],'inv':[row]})
        statement['docdata']['b2b'].append({'ctin':gstin('33',900),'inv':[{'inum':'STATEMENT-ONLY-001','dt':'20-10-2026','txval':'1000.00','cgst':'90.00','sgst':'90.00','igst':'0.00','cess':'0.00','itcavl':'Y'}]})
        (ROOT/'gstr2b'/f'{b["key"]}.json').write_text(json.dumps(statement,indent=2),encoding='utf-8', newline='\n')
        duplicate=deepcopy(statement);duplicate['docdata']['b2b'][0]['inv']*=2
        (ROOT/'edge_cases'/f'{b["key"]}_duplicate_statement.json').write_text(json.dumps(duplicate,indent=2),encoding='utf-8', newline='\n')
    base=manifest['invoices'][0]
    cases=[('duplicate_invoice','business duplicate'),('renamed_same_file','exact file duplicate'),('invalid_gstin','block confirmation'),('missing_supplier_gstin','block confirmation'),('missing_date','block confirmation'),('missing_number','block confirmation'),('incorrect_total','block confirmation'),('incorrect_tax','block confirmation'),('invalid_tax_heads','block confirmation'),('low_quality','manual review; extraction may fail'),('rotated','manual review'),('multi_page','one invoice, two pages'),('credit_note','signed reduction after validated confirmation'),('debit_note','signed increase after validated confirmation'),('multiple_rates','validate each line separately'),('incomplete_supplier','manual correction required'),('different_financial_year','distinct FY; not automatically duplicate'),('unsupported_upload','reject unsupported content')]
    for name,expected in cases:
        inv=deepcopy(base);ext='pdf';pages=1
        if name not in ['duplicate_invoice','renamed_same_file','different_financial_year']:
            inv['invoice_number']='EDGE-'+name.replace('_','-')[:11]
        if name=='invalid_gstin':inv['supplier']['gstin']='INVALID-GSTIN'
        if name=='missing_supplier_gstin':inv['supplier']['gstin']=''
        if name=='missing_date':inv['invoice_date']=''
        if name=='missing_number':inv['invoice_number']=''
        if name=='incorrect_total':inv['totals']['invoice_total']='1.00'
        if name=='incorrect_tax':inv['line_items'][0]['igst']='1.00'
        if name=='invalid_tax_heads':inv['line_items'][0]['cgst']='5.00'
        if name=='incomplete_supplier':inv['supplier']['legal_name']='Not provided';inv['supplier']['principal_address']='Not provided'
        if name in ['credit_note','debit_note']:inv['invoice_type']=name;inv['invoice_number']='EDGE-'+name.replace('_','-').upper()[:11]
        if name=='different_financial_year':inv['invoice_date']='2025-10-01'
        image=render(inv)
        if name=='low_quality':image=image.resize((310,438)).filter(ImageFilter.GaussianBlur(1.5)).resize((1240,1754));ext='jpg'
        if name=='rotated':image=image.rotate(90,expand=True);ext='jpg'
        if name=='multi_page':pages=2
        if name=='unsupported_upload':ext='txt'
        path=ROOT/'edge_cases'/(name+'.'+ext)
        if name in ['duplicate_invoice','renamed_same_file']:path.write_bytes((ROOT/base['files'][0]).read_bytes())
        elif ext=='txt':path.write_text('Synthetic unsupported upload - not an invoice',encoding='utf-8', newline='\n')
        elif ext=='pdf':pdf(image,path,pages)
        else:image.save(path)
        manifest['edge_cases'].append({'file':str(path.relative_to(ROOT)).replace('\\','/'),'expected':expected})
    assumptions=wb.create_sheet('Assumptions');assumptions.append(['No dashboard data is provisioned. These are expected results ONLY after manual upload/review/confirmation.']);assumptions.append(['Tax rates and HSN/SAC are synthetic illustrative assumptions, not verified legal classifications.']);assumptions.append(['All ITC starts pending review. Matching alone does not approve ITC. B2B parser does not support credit-note sections.']);assumptions.append(['Choose one format per invoice. Alternate formats are deliberately the same business invoice.'])
    for tab in wb:
        tab.freeze_panes='A2'
        for col in tab.columns:tab.column_dimensions[col[0].column_letter].width=25
    wb.save(ROOT/'EXPECTED_RESULTS.xlsx')
    manifest['file_hashes']={str(p.relative_to(ROOT)).replace('\\','/'):hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(ROOT.rglob('*')) if p.is_file() and p.name not in {'TEST_MANIFEST.json','README.md'}}
    (ROOT/'TEST_MANIFEST.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8', newline='\n')
    (ROOT/'README.md').write_text('# Nature Labs synthetic upload pack\n\n28 distinct invoices (12 manufacturing, 8 trading, 8 services), each in PDF, PNG and JPEG. Upload ONE representation per invoice. All are fictional and marked NOT VALID FOR TAX FILING. No real bank details are used. GSTIN-like identifiers satisfy checksum but are not verified taxpayer identities. Do not use outside the isolated V2 demo.\n\nThe environment is provisioned with profiles ONLY. Upload files yourself, review extracted values and confirm valid invoices before expecting financial dashboard/report figures. Select October 2026. Expected totals are in EXPECTED_RESULTS.xlsx; monetary calculations use Decimal. ITC remains pending until evidence-based review. Test rates (12%, 18%) and HSN/SAC assignments are assumptions, not current-law advice.\n\nGSTR-2B fixtures contain supported B2B matches, amount/tax-head differences and statement-only records. Other purchases are books-only after import. Duplicate statements in edge_cases must be rejected; credit-note GSTR sections are unsupported, not simulated. Imported statements never approve ITC. Edge-case outcomes are in TEST_MANIFEST.json. Credit notes use positive face amounts and negative signed financial contributions.\n\nUse python scripts/generate_gst_easy_v2_demo.py to regenerate with Pillow, ReportLab and openpyxl. Generator has no cloud or application API calls. Logos are fictional programmatic monograms.\n',encoding='utf-8', newline='\n')
    print('Generated',len(manifest['invoices']),'invoices in three formats; no application records written.')


if __name__=='__main__':main()
