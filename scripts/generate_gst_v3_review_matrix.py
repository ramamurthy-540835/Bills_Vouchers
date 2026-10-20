"""Generate the requested professional-review workbook without legal assertions."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

SCENARIOS = [
('Regular monthly filer: sales and purchases','Regular','GSTR-1 / GSTR-3B','CGST 37,39'),
('No outward sales, inward purchases','Regular','GSTR-3B / GSTR-2B','CGST 16,39'),
('Nil-return-eligible taxpayer','Applicable filer','Applicable return','CGST 39'),
('Sales with no eligible ITC','Regular','GSTR-3B','CGST 16,17,49'),
('ITC greater than liability','Regular','GSTR-3B / credit ledger','CGST 49,49A,49B'),
('Liability greater than ITC','Regular','GSTR-3B / cash ledger','CGST 49,49A,49B'),
('QRMP taxpayer','Eligible regular','IFF / GSTR-1 / GSTR-3B / PMT-06','CGST 37,39'),
('Composition taxpayer','Eligible composition','CMP-08 / GSTR-4','CGST 10'),
('Manufacturing business','Regular manufacturer','GSTR-1 / GSTR-3B','CGST 7,9,16,17'),
('Wholesale trader','Regular trader','GSTR-1 / GSTR-3B','CGST 7,9,16'),
('Retail trader','Applicable taxpayer','GSTR-1 / GSTR-3B','CGST 7,9,31'),
('Service provider','Applicable taxpayer','GSTR-1 / GSTR-3B','CGST 13,15 / IGST 12,13'),
('Export-oriented business','Exporter','GSTR-1 / GSTR-3B / RFD-01','IGST 16 / CGST 54'),
('SEZ supplier','SEZ supply','GSTR-1 / GSTR-3B','IGST 16'),
('Goods importer','Importer','GSTR-3B / bill of entry','IGST 5 / CGST 16'),
('Services importer','Importer','GSTR-3B RCM','IGST 5,13'),
('Reverse-charge transactions','Notified recipient','GSTR-3B','CGST 9 / IGST 5'),
('Credit/debit notes','Applicable supplier','GSTR-1 / GSTR-3B','CGST 34'),
('Invoice amendments','Applicable filer','GSTR-1 / GSTR-1A','CGST 37'),
('Blocked ITC','Regular','GSTR-3B table 4','CGST 17(5)'),
('Mixed taxable and exempt supplies','Mixed supplies','GSTR-3B table 4','CGST 17 / Rules 42,43'),
('Capital goods','Regular','GSTR-3B table 4','CGST 16,17 / Rule 43'),
('Common-credit reversals','Regular','GSTR-3B table 4','CGST 17 / Rules 42,43'),
('Unpaid supplier reversal conditions','Regular','GSTR-3B table 4','CGST 16 / Rule 37'),
('Purchases missing from GSTR-2B','Regular','GSTR-2B / GSTR-3B','CGST 16 / Rule 36'),
('IMS pending/rejected documents','Eligible IMS user','IMS / GSTR-2B','Current GST portal IMS guidance'),
('E-commerce seller','Seller','GSTR-1 / GSTR-3B','CGST 9(5),52'),
('E-commerce operator','Operator','GSTR-8 / applicable returns','CGST 9(5),52'),
('GST TDS deductor','TDS deductor','GSTR-7 / GSTR-7A','CGST 51'),
('Input Service Distributor','ISD','GSTR-6','CGST 20 / Rule 39'),
('Non-resident taxpayer','Non-resident taxable','GSTR-5','CGST 24,27,39'),
('OIDAR taxpayer','Applicable OIDAR','GSTR-5A','IGST 14'),
('UIN holder','UIN','GSTR-11','CGST 25,55'),
('Business with multiple GSTINs','Multiple registrations','Per-registration returns','CGST 25'),
('Late return filing','Applicable filer','Interest / late-fee working paper','CGST 47,50'),
('Missed invoices','Applicable filer','Permitted subsequent amendment','CGST 37,39'),
('Excess tax payment','Eligible taxpayer','Refund / ledger','CGST 54'),
('GST refund application','Eligible claimant','RFD-01','CGST 54'),
('GST notice','Noticed taxpayer','REG / ASMT / DRC / APL','Applicable notice-specific provisions'),
('Registration cancellation','Registered taxpayer','REG cancellation','CGST 29'),
('Final return','Applicable cancelled registration','GSTR-10','CGST 45'),
('Annual GSTR-9/GSTR-9C','Applicable annual filer','GSTR-9 / GSTR-9C','CGST 44'),
('E-invoicing applicability','Notified taxpayer','IRP schema','CGST 31 / Rule 48'),
('E-way bill applicability','Applicable movement','E-way bill','CGST 68 / Rule 138'),
('Branch/stock transfers','Distinct persons','GSTR-1 / GSTR-3B','CGST 25 / Schedule I'),
('Job work','Principal / job worker','Applicable job-work records','CGST 19,143'),
('Special place-of-supply cases','Relevant supply','Applicable return','IGST 10–13'),
('Period-specific GST rate changes','Effective-date taxpayer','Applicable return','CGST 14 / rate notifications'),
('Filing locks and post-filing corrections','Filed period','Separate amendment workflow','CGST 37,39'),
('Incomplete/unreliable invoice evidence','All','Review queue','CGST 16,31,35'),
]


def main():
    book=Workbook()
    sheet=book.active
    sheet.title='50 scenario review'
    sheet.append(['ID','Scenario','Eligibility to verify','Form/table candidate','Legal reference candidate',
                  'Official source','Required evidence','Calculation / independent result',
                  'Validation / exception','Approval','Synthetic fixture','Engineering evidence',
                  'Test result','Legal review status'])
    for index,(name,taxpayer,form,legal) in enumerate(SCENARIOS,1):
        engineering = ('Existing canonical report tests; new declared-rate comparison tests'
                       if index in {1,2,4,5,6,9,10,11,12,18,20,25,49,50}
                       else 'No scenario-complete implementation acceptance established')
        sheet.append([f'{index:02}',name,taxpayer,form,legal,
            'https://www.indiacode.nic.in/ ; https://tutorial.gst.gov.in/ ; https://cbic-gst.gov.in/',
            f'Period, registration category, source invoices and scenario-specific evidence for: {name}',
            'Independent specialist calculation required; no statutory amount invented',
            'Validate applicability, effective date, evidence, arithmetic and immutable period; unresolved facts block approval',
            'Preparer plus independent authorized reviewer; government submission disabled',
            'Scenario fixture pending professional specification',
            engineering,'NOT_EXECUTED as complete legal scenario','PENDING_AUTHORIZED_SPECIALIST_REVIEW'])
    sheet.freeze_panes='C2';sheet.auto_filter.ref=sheet.dimensions
    for cell in sheet[1]:
        cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='187653')
    for row in sheet.iter_rows(min_row=2):
        for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
    for column in sheet.columns:sheet.column_dimensions[column[0].column_letter].width=32
    sheet.column_dimensions['A'].width=8
    notes=book.create_sheet('Review boundary')
    notes.append(['Status','Meaning'])
    notes.append(['Not statutory certification','References are candidates for effective-date review, not verified legal conclusions.'])
    notes.append(['Engineering tests','Local tests do not certify all 50 scenarios, official exports, or government acceptance.'])
    notes.append(['Direct filing','Disabled pending approved GSP credentials, taxpayer consent and verified schemas.'])
    output=Path('GST_EASY_50_SCENARIO_TEST_MATRIX.xlsx')
    book.save(output)
    print(f'Created {output}: {len(SCENARIOS)} scenarios, all formal legal acceptance pending.')


if __name__=='__main__':
    main()
