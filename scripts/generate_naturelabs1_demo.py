"""Deterministic synthetic evidence pack; no cloud access or application writes."""
import hashlib
import json
from copy import deepcopy
from decimal import Decimal as D
from pathlib import Path

from PIL import Image, ImageDraw
from openpyxl import Workbook
from generate_gst_easy_v2_demo import gstin, money, render, pdf, font, HEADS

ROOT = Path('demo bills/naturelabs1')
PERIOD = '2026-10'
BUSINESS = dict(
    key='naturelabs1', legal_name='Nature Labs Synthetic Manufacturing Pvt Ltd',
    trade_name='Nature Labs Industrial Manufacturing', constitution='Company',
    business_nature='Manufacturing industrial electrical panels, electronic components, sensors and machinery equipment',
    state_code='33', city='Chennai', pin='600032',
    principal_address='No. 25, Fictional Industrial Estate, Demo Technology Road, Chennai, Tamil Nadu - 600032',
    contact_email='manufacturing-demo@example.test', contact_phone='+919000000201',
    website_url='https://naturelabs-manufacturing.example.test',
    gstin=gstin('33',201), pan=gstin('33',201)[2:12],
    registration_type='Regular', filing_frequency='monthly',
    is_test=True, gstin_verified=False, provenance='SYNTHETIC_TEST_ONLY')
PURCHASES = ['Steel raw material','Copper wire','Electronic components','Industrial sensors',
             'Packaging materials','Machinery spare parts','Factory maintenance services',
             'Electrical equipment','Interstate raw material','Purchase credit note']
SALES = ['Industrial control panels','Electrical distribution boards','Sensor assemblies',
         'Industrial wiring equipment','Fabricated steel components','Machinery accessories',
         'Industrial automation parts','Electronic control units','Interstate finished goods','Sales credit note']


def make_invoice(index, direction, status='confirmed'):
    n=index+1
    state='29' if index in {0,5,8} else '33'
    counter=dict(legal_name=f'Fictional {"Supplier" if direction=="purchase" else "Customer"} {n:02d}',
                 gstin=gstin(state,300+index+(30 if direction=='sale' else 0)), state_code=state,
                 principal_address=f'Unit {n}, Synthetic Industrial Road, Demo City (fictional)')
    supplier, recipient=(counter,BUSINESS) if direction=='purchase' else (BUSINESS,counter)
    description=(PURCHASES if direction=='purchase' else SALES)[index%10]
    items=[]
    for j in range(2):
        qty=D(2+j); price=D(1200+n*193+j*71)*(D(2) if direction=='sale' else D(1))
        taxable=qty*price
        # Explicit zero-rate test fixture, not a legal classification of packaging.
        rate=D(0 if direction=='purchase' and index==4 else 18 if j==0 else 12)
        tax=(taxable*rate/100).quantize(D('.01'))
        items.append(dict(description=description+(' - accessory' if j else ''),hsn_sac='853710',
                          quantity=str(qty),unit='NOS',unit_price=money(price),taxable_value=money(taxable),
                          gst_rate=str(rate),cgst=money(0 if state=='29' else tax/2),
                          sgst=money(0 if state=='29' else tax/2),igst=money(tax if state=='29' else 0),cess='0.00'))
    totals={h:money(sum((D(line[h]) for line in items),D(0))) for h in HEADS[:-1]}
    totals['invoice_total']=money(sum(map(D,totals.values()),D(0)))
    return dict(business='naturelabs1',invoice_number=f'NL1-{direction[0].upper()}-{n:03d}',
                invoice_date=f'2026-10-{n:02d}',invoice_type='credit_note' if index==9 else 'tax_invoice',
                direction=direction,supplier=deepcopy(supplier),recipient=deepcopy(recipient),
                place_of_supply_state_code=recipient['state_code'],line_items=items,totals=totals,
                expected_status=status,provenance='SYNTHETIC_TEST_ONLY')


def main():
    ROOT.mkdir(parents=True,exist_ok=True)
    docs=[make_invoice(i,d) for d in ['purchase','sale'] for i in range(10)]
    docs += [make_invoice(10,'purchase','needs_review'),make_invoice(11,'sale','needs_review'),make_invoice(12,'purchase','rejected')]
    for inv in docs:
        image=render(inv); inv['files']=[]
        for ext in ['pdf','png','jpg']:
            path=ROOT/(inv['invoice_number']+'.'+ext)
            if ext=='pdf':pdf(image,path)
            else:image.save(path)
            inv['files'].append(path.name)
    logo=Image.new('RGB',(320,320),'#187653');draw=ImageDraw.Draw(logo)
    draw.rounded_rectangle((20,20,300,300),radius=32,outline='white',width=4)
    draw.text((65,65),'NL',font=font(110,True),fill='white')
    draw.text((40,225),'INDUSTRIAL',font=font(32,True),fill='white');logo.save(ROOT/'logo.png')
    statement={'gstin':BUSINESS['gstin'],'rtnprd':'102026','docdata':{'b2b':[]}}
    for index in [0,1,2,3,5,6,7]:
        inv=docs[index]; t=inv['totals']
        row={'inum':inv['invoice_number'],'dt':'-'.join(reversed(inv['invoice_date'].split('-'))),'itcavl':'Y',
             **{out:t[src] for out,src in [('txval','taxable_value'),('cgst','cgst'),('sgst','sgst'),('igst','igst'),('cess','cess')]}}
        if index==1:row['cgst']=money(D(row['cgst'])+25)
        if index==2:row['igst']=money(D(row['cgst'])+D(row['sgst']));row['cgst']=row['sgst']='0.00'
        statement['docdata']['b2b'].append({'ctin':inv['supplier']['gstin'],'inv':[row]})
    statement['docdata']['b2b'].append({'ctin':gstin('33',499),'inv':[dict(inum='NL1-ONLY-001',dt='15-10-2026',txval='2000.00',cgst='180.00',sgst='180.00',igst='0.00',cess='0.00',itcavl='Y')]})
    (ROOT/'gstr2b.json').write_text(json.dumps(statement,indent=2),encoding='utf-8',newline='\n')
    duplicate=deepcopy(statement);duplicate['docdata']['b2b'][0]['inv']*=2
    (ROOT/'gstr2b-duplicate-invalid.json').write_text(json.dumps(duplicate,indent=2),encoding='utf-8',newline='\n')
    # Independent arithmetic over fixture face values, never calls dashboard/report engines.
    totals={direction:{h:money(sum((D(i['totals'][h])*(-1 if i['invoice_type']=='credit_note' else 1)
                     for i in docs if i['direction']==direction and i['expected_status']=='confirmed'),D(0)))
                     for h in HEADS} for direction in ['purchase','sale']}
    credit={h:max(D(0),sum((D(docs[i]['totals'][h]) for i in [0,3,5,6,7]),D(0))-D(docs[9]['totals'][h])) for h in ['cgst','sgst','igst','cess']}
    output=sum((D(totals['sale'][h]) for h in credit),D(0)); eligible=sum(credit.values(),D(0))
    expected=dict(eligible_credit=money(eligible),output_tax=money(output),cash_required=money(output-eligible),
                  non_gst=docs[4]['totals']['invoice_total'],confirmed_sales=totals['sale']['invoice_total'],
                  confirmed_purchases=totals['purchase']['invoice_total'],pending_reviews=2,
                  reconciliation_mismatches=3,unmatched_invoices=3)
    manifest=dict(provenance='SYNTHETIC_TEST_ONLY',period=PERIOD,business=BUSINESS,invoices=docs,
                  expected_signed_totals=totals,expected_kpis=expected,
                  assumptions=['0%, 12%, 18% illustrative test rates; not legally verified HSN classifications.',
                               'Zero-rate packaging fixture exercises non-creditable reporting only.',
                               'Credit notes have positive face values and negative signed contributions.',
                               'Five matched eligible purchases less the purchase credit note form eligible credit.',
                               'GSTR-2B B2B only. Duplicate fixture must be rejected; credit-note sections unsupported.',
                               'Reconciliation exceptions: two amount/head differences plus one unsupported credit note; three unmatched records.',
                               'GSTIN-like values are fictional by construction; no government lookup or ownership claim.'])
    wb=Workbook();ws=wb.active;ws.title='Invoices';ws.append(['Invoice','Direction','Type','Status',*HEADS])
    for i in docs:ws.append([i['invoice_number'],i['direction'],i['invoice_type'],i['expected_status'],*[i['totals'][h] for h in HEADS]])
    ws=wb.create_sheet('Expected KPIs');ws.append(['Metric','Expected']);[ws.append([k,v]) for k,v in expected.items()]
    ws=wb.create_sheet('Signed totals');ws.append(['Direction',*HEADS]);[ws.append([d,*[v[h] for h in HEADS]]) for d,v in totals.items()]
    ws=wb.create_sheet('Assumptions');[ws.append([s]) for s in manifest['assumptions']]
    wb.save(ROOT/'EXPECTED_RESULTS.xlsx')
    manifest['file_hashes']={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in ROOT.iterdir() if p.is_file() and p.name not in {'TEST_MANIFEST.json','README.md'}}
    (ROOT/'TEST_MANIFEST.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8',newline='\n')
    (ROOT/'README.md').write_text('# Nature Labs 1 synthetic evidence\n\n23 distinct invoices: 10 confirmed purchases, 10 confirmed sales (including one credit note each), two pending and one rejected example. Each has PDF, PNG and JPEG representations. Upload only one format per identity. Original documents are visibly fictional and not valid for filing.\n\nOctober 2026. EXPECTED_RESULTS.xlsx and TEST_MANIFEST.json contain independently calculated expectations, not application data. Provisioning and live acceptance must be recorded separately. Rates and zero-rate treatment are test assumptions, not tax advice. Duplicate statement fixture must be rejected. No government verification is asserted.\n\nRegenerate with scripts/generate_naturelabs1_demo.py. This generator makes no cloud calls and does not modify the previous demo pack or accounts.\n',encoding='utf-8',newline='\n')
    print(f'Generated {len(docs)} synthetic invoices and independent expectations in {ROOT}')


if __name__=='__main__':main()
