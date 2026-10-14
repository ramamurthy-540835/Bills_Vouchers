from types import SimpleNamespace
import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from app.db import get_db
from app.routes import current_user
from app.config import get_settings
from app.routers import gst_bills


@pytest.fixture
def api(monkeypatch,invoice,profile):
    monkeypatch.setenv('GST_BILLS_ENABLED','true');get_settings.cache_clear()
    repo=SimpleNamespace()
    user=SimpleNamespace(id='u',role='viewer')
    items={invoice.invoice_id:invoice}
    deleted={}
    class Store:
        def assert_writable(self,period):pass
        def __init__(self,repo,client_id):self.client_id=client_id
        def list(self,*args):return [i for i in items.values() if i.client_id==self.client_id]
        def matching_rows(self,period):return []
        def get(self,id):
            inv=items.get(id)
            if not inv or inv.client_id!=self.client_id:raise HTTPException(404,'Bill not found.')
            return inv
        def save(self,inv,*args,**kwargs):items[inv.invoice_id]=inv
        def delete(self,inv,*args):deleted[inv.invoice_id]=inv;del items[inv.invoice_id]
        def deleted_invoice(self,id):
            if id not in deleted or deleted[id].client_id!=self.client_id:raise HTTPException(404)
            return deleted[id]
    monkeypatch.setattr(gst_bills,'BillStore',Store)
    monkeypatch.setattr(gst_bills,'active_client',lambda *a:SimpleNamespace(id='a'))
    monkeypatch.setattr(gst_bills,'profile_for',lambda *a:profile)
    import app.services.gst.account as account
    monkeypatch.setattr(account,'FinanceRepository',lambda _:SimpleNamespace(can_access_client=lambda uid,cid:cid=='a'))
    app=FastAPI();app.include_router(gst_bills.router)
    app.dependency_overrides[get_db]=lambda:repo
    app.dependency_overrides[current_user]=lambda:user
    yield TestClient(app),user,items
    get_settings.cache_clear()


def test_viewer_read_export_but_no_writes(api,invoice):
    client,user,_=api
    assert client.get('/api/gst/bills?period=2026-09').status_code==200
    assert client.get('/api/gst/bills/export?period=2026-09&format=csv').status_code==200
    assert client.get('/api/gst/bills/template.xlsx').status_code==200
    assert client.post('/api/gst/bills/upload?period=2026-09',files={'files':('bill.csv',b'csv','text/csv')}).status_code==403
    assert client.post(f'/api/gst/bills/{invoice.invoice_id}/confirm').status_code==403
    assert client.patch(f'/api/gst/bills/{invoice.invoice_id}',json={'invoice_number':'CHANGE'}).status_code==403
    assert client.delete(f'/api/gst/bills/{invoice.invoice_id}').status_code==403
    assert client.post('/api/gst/bills/upload/other-client-job/retry').status_code==403


def test_reconciliation_is_scoped_and_export_retains_evidence(api, invoice):
    from io import BytesIO
    from openpyxl import load_workbook
    client, user, items = api
    invoice.status = 'confirmed'
    other = invoice.model_copy(deep=True)
    other.invoice_id = 'private-other-client'
    other.client_id = 'b'
    items[other.invoice_id] = other
    response = client.get('/api/gst/bills/reconciliation?period=2026-09')
    assert response.status_code == 200
    assert response.json()['counts'] == {'missing_statement': 1}
    assert other.invoice_id not in response.text
    assert client.get('/api/gst/bills/reconciliation?period=bad').status_code == 422
    response = client.get('/api/gst/bills/working-papers?period=2026-09')
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert workbook['Reconciliation evidence']['A3'].value == invoice.invoice_id
    assert workbook['Reconciliation evidence']['E3'].value == 'missing_statement'
    assert 'Statement candidates' in workbook.sheetnames


def test_delete_checks_version_and_removes_bill(api,invoice):
    client,user,items=api;user.role='client'
    url=f'/api/gst/bills/{invoice.invoice_id}'
    assert client.delete(url,headers={'If-Match':'999'}).status_code==409
    assert invoice.invoice_id in items
    assert client.delete('/api/gst/bills/another-client').status_code==404
    assert client.delete(url,headers={'If-Match':str(invoice.version)}).status_code==200
    assert client.get(url).status_code==404
    assert client.get('/api/gst/bills?period=2026-09').json()['invoices']==[]


def test_cross_client_id_is_not_found(api,invoice):
    client,user,items=api
    other=invoice.model_copy(deep=True);other.invoice_id='other';other.client_id='b';items['other']=other
    for role in ('viewer','client','admin','tax_admin'):
        user.role=role
        assert client.get('/api/gst/bills/other').status_code==404
        assert client.get('/api/gst/bills/other/original').status_code==404
        assert 'other' not in client.get('/api/gst/bills?period=2026-09').text
    assert client.post('/api/gst/bills/other/confirm').status_code==404
    assert client.patch('/api/gst/bills/other',json={'invoice_number':'X'}).status_code==404


def test_client_can_confirm_but_not_change_tenant(api,invoice):
    client,user,items=api;user.role='client'
    assert client.patch(f'/api/gst/bills/{invoice.invoice_id}',json={'client_id':'b'}).status_code==422
    response=client.post(f'/api/gst/bills/{invoice.invoice_id}/confirm')
    assert response.status_code==200 and response.json()['status']=='confirmed'
    assert isinstance(response.json()['totals']['invoice_total'],str)


def test_sample_export_blocked(api,invoice):
    client,user,items=api;invoice.sample=True
    response=client.get('/api/gst/bills/export?period=2026-09')
    assert response.status_code==403 and response.json()['detail']['code']=='sample_data_blocked'


def test_test_client_can_review_but_cannot_export(api,monkeypatch):
    client,user,_=api
    monkeypatch.setattr(gst_bills,'active_client',lambda *a:SimpleNamespace(id='a',is_test=True))
    assert client.get('/api/gst/bills?period=2026-09').status_code==200
    assert client.get('/api/gst/bills/template.xlsx').status_code==200
    for path in ('export?period=2026-09&format=csv','export?period=2026-09&format=xlsx','working-papers?period=2026-09'):
        response=client.get('/api/gst/bills/'+path)
        assert response.status_code==403 and response.json()['detail']['code']=='sample_data_blocked'


def test_missing_profile_prevents_confirmation(api,invoice,monkeypatch):
    client,user,_=api;user.role='client'
    monkeypatch.setattr(gst_bills,'profile_for',lambda *a:{})
    assert client.post(f'/api/gst/bills/{invoice.invoice_id}/confirm').status_code==409


def test_rejection_requires_reason_and_preserves_original(api,invoice):
    client,user,items=api;user.role='client'
    url=f'/api/gst/bills/{invoice.invoice_id}/reject'
    for body in ({}, {'reason':' '}, {'reason':123}, {'reason':'x'*1001}):
        assert client.post(url,json=body).status_code==422
        assert items[invoice.invoice_id].status==invoice.status
    response=client.post(url,json={'reason':' Supplier correction required. '},headers={'If-Match':str(invoice.version)})
    assert response.status_code==200
    assert response.json()['rejection_reason']=='Supplier correction required.'
    assert response.json()['version']==invoice.version+1
    assert invoice.status!='rejected'
    assert client.post(url,json={'reason':'Old edit'},headers={'If-Match':str(invoice.version)}).status_code==409


def test_restore_requires_write_permission_and_returns_to_review(api,invoice):
    client,user,items=api;user.role='client'
    invoice.status='confirmed'
    assert client.delete(f'/api/gst/bills/{invoice.invoice_id}').status_code==200
    user.role='viewer'
    url=f'/api/gst/bills/{invoice.invoice_id}/restore'
    assert client.post(url).status_code==403
    user.role='client'
    assert client.post('/api/gst/bills/other-client/restore').status_code==404
    assert client.post(url,headers={'If-Match':'999'}).status_code==409
    result=client.post(url,headers={'If-Match':str(invoice.version)})
    assert result.status_code==200
    assert result.json()['status']=='needs_review'
    assert result.json()['version']==invoice.version+1


def test_line_sum_uses_decimal_and_does_not_write(api,invoice):
    client,user,items=api
    body={'line_items':[{'taxable_value':'0.10'},{'taxable_value':'0.20','igst':'0.05'}]}
    assert client.post('/api/gst/bills/sum-lines',json=body).status_code==403
    user.role='client'
    original=invoice.model_dump_json()
    result=client.post('/api/gst/bills/sum-lines',json=body)
    assert result.status_code==200
    assert result.json()['taxable_value']=='0.30'
    assert result.json()['invoice_total']=='0.35'
    assert items[invoice.invoice_id].model_dump_json()==original
    body['line_items'][0]['taxable_value']=0.1
    assert client.post('/api/gst/bills/sum-lines',json=body).status_code==422


def test_stale_edit_and_payment_authority(api,invoice):
    client,user,_=api;user.role='client'
    assert client.patch(f'/api/gst/bills/{invoice.invoice_id}',headers={'If-Match':'999'},json={'invoice_number':'NEW'}).status_code==409
    assert client.patch(f'/api/gst/bills/{invoice.invoice_id}',json={'reverse_charge_paid_reference':'REF123'}).status_code==403
    assert client.patch(f'/api/gst/bills/{invoice.invoice_id}',json={'totals':{'taxable_value':'bad'}}).status_code==422
