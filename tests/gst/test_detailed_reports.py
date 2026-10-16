from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from app.gst import reports
from tests.gst import persisted_reports


@pytest.fixture
def persisted_store(tmp_path,monkeypatch):
    path=tmp_path/'reports.sqlite'
    persisted_reports.seed(path)
    store=SimpleNamespace(client_id='browser-client')
    store.list=lambda period:[inv for inv in persisted_reports.invoices(path,store.client_id) if str(inv.invoice_date)[:7]==period]
    store.matching_rows=lambda period:persisted_reports.statements(path,store.client_id,period)
    store.statement_history=lambda period:[]
    store.query=lambda *args:[]
    store.table=lambda *args:'synthetic.table'
    monkeypatch.setattr(reports.returns,'fingerprint',lambda *args:'a'*64)
    return store


def build(store,kind,filters=None,period='2026-09'):
    return reports.build(store,period,'month',kind,lambda _: {'business_nature':'other','legal_name':'Synthetic persisted test'},filters)


@pytest.mark.parametrize('kind,count',[('credit',4),('liability',2),('cash',16),('without_itc',3),('pending',1),('mismatches',1),('unmatched',2),('sales',2),('purchases',4)])
def test_all_nine_persisted_projections(persisted_store,kind,count):
    result=build(persisted_store,kind)
    assert result['row_count']==count
    assert result['client_id']=='browser-client'
    assert all(len(row)==len(result['columns']) for row in result['rows'])
    for format in ('csv','xlsx'):
        content,mime=reports.render(result,format)
        assert content and mime
    if kind=='sales':assert result['summary']['Invoice total']=='5310.00'
    if kind=='credit':assert result['monthly_summary'][0]['totals']['eligible_credit']=='180.00'
    if kind=='cash':
        assert result['monthly_cash_scenarios'][0]=={'period':'2026-09','recorded':'0.00','matched':'270.00','eligible':'630.00','noCredit':'810.00'}
        assert sum(Decimal(row[result['columns'].index('Credit applied to liability')]) for row in result['rows'] if row[1]=='eligible')==Decimal('180')


def test_filters_exports_and_period_are_consistent(persisted_store):
    result=build(persisted_store,'sales',{'invoice_type':'credit_note','q':'CREDIT','date_from':'2026-09-01','date_to':'2026-09-30','sort':'-Invoice total'})
    assert result['row_count']==1 and result['summary']['Invoice total']=='-590.00'
    book=load_workbook(BytesIO(reports.render(result,'xlsx')[0]))
    assert book['Report'].max_row==2
    assert book['Source invoices'].max_row==2
    assert build(persisted_store,'sales',period='2026-08')['row_count']==0
    with pytest.raises(HTTPException):build(persisted_store,'sales',{'date_from':'bad'})
    with pytest.raises(HTTPException):build(persisted_store,'sales',{'sort':'arbitrary_column'})


def test_missing_statement_is_dependency_not_unmatched(persisted_store):
    persisted_store.matching_rows=lambda period:[]
    result=build(persisted_store,'unmatched')
    assert result['rows']==[] and 'Statement not imported' in result['warnings'][0]
    persisted_store.statement_history=lambda period:[{'record_count':0}]
    result=build(persisted_store,'unmatched')
    assert result['row_count']==4
    assert all(row[6]=='missing_statement' for row in result['rows'])


def test_business_switch_changes_persisted_rows_without_mixing(persisted_store):
    first=build(persisted_store,'sales')
    persisted_store.client_id='browser-readonly'
    second=build(persisted_store,'sales')
    assert second['summary']['Invoice total']=='10620.00'
    assert not ({row[1] for row in first['rows']} & {row[1] for row in second['rows']})
