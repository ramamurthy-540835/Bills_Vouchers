import csv
from copy import deepcopy
from io import BytesIO, StringIO
from unittest.mock import Mock

import pdfplumber
import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.gst import reports


@pytest.fixture
def report_store(invoice, monkeypatch):
    invoice.status = 'confirmed'
    store = Mock(client_id='a')
    store.list.side_effect = lambda month: [deepcopy(invoice)] if month == '2026-09' else []
    store.matching_rows.return_value = []
    store.query.return_value = []
    monkeypatch.setattr(reports.returns, 'fingerprint', lambda *args: 'a'*64)
    return store


def test_periods_cover_indian_financial_year_and_quarter():
    assert reports.periods('2026-01','quarter') == ['2026-01','2026-02','2026-03']
    assert reports.periods('2026-03','financial_year') == [f'2025-{month:02d}' for month in range(4,13)]+['2026-01','2026-02','2026-03']
    assert reports.periods('2026-04','financial_year')[0] == '2026-04'
    with pytest.raises(HTTPException):
        reports.periods('2026-09','anything')


def test_confirmed_registers_are_separate_and_credit_notes_signed(report_store, invoice, profile):
    sale = invoice.model_copy(deep=True);sale.direction='sale';sale.invoice_id='sale'
    rejected = invoice.model_copy(deep=True);rejected.status='rejected';rejected.invoice_id='rejected'
    note = invoice.model_copy(deep=True);note.invoice_type='credit_note';note.invoice_id='note'
    report_store.list.return_value = [invoice,sale,rejected,note]
    report_store.list.side_effect = None
    purchase = reports.build(report_store,'2026-09','month','purchases',lambda _:profile)
    assert purchase['row_count'] == 2
    assert next(row[-1] for row in purchase['rows'] if row[1]=='note') == '-1180.00'
    sales = reports.build(report_store,'2026-09','month','sales',lambda _:profile)
    assert sales['row_count']==1 and sales['rows'][0][1]=='sale'
    listing = reports.build(report_store,'2026-09','month','invoices',lambda _:profile)
    assert listing['row_count']==4


def test_xlsx_has_provenance_and_formula_strings_are_not_executable(report_store, invoice, profile):
    invoice.invoice_number='  =HYPERLINK("bad")'
    invoice.supplier_legal_name='@SUM(1,1)'
    result=reports.build(report_store,'2026-09','month','purchases',lambda _:profile)
    content,mime=reports.render(result,'xlsx')
    book=load_workbook(BytesIO(content))
    assert set(book.sheetnames)=={'Report','Provenance','Source invoices'}
    assert book['Report']['J2'].value.startswith("'")
    assert book['Report']['G2'].data_type=='s'
    assert book['Source invoices']['B2'].value==invoice.invoice_id
    assert 'raw_ref' not in str(list(book['Provenance'].values))
    content,_=reports.render(result,'csv')
    rows=list(csv.reader(StringIO(content.decode('utf-8-sig'))))
    assert any(row and row[0]=='Report SHA256' for row in rows)
    assert any(len(row)>9 and row[9].startswith("'  =") for row in rows)


def test_tax_pdf_and_quarter_preserve_monthly_scenarios(report_store, profile):
    result=reports.build(report_store,'2026-09','quarter','tax',lambda _:profile)
    assert result['periods']==['2026-07','2026-08','2026-09']
    assert result['row_count']==12 and len(result['monthly_cash_scenarios'])==3
    assert set(result['monthly_cash_scenarios'][0])=={'period','recorded','matched','eligible','noCredit'}
    content,mime=reports.render(result,'pdf')
    assert mime=='application/pdf' and content.startswith(b'%PDF')
    with pdfplumber.open(BytesIO(content)) as document:
        text='\n'.join(page.extract_text() or '' for page in document.pages)
    assert '2026-07' in text and '2026-09' in text and 'not a GST filing export' in text
    assert result['report_hash'] in text


def test_changed_source_denies_report(report_store, profile, monkeypatch):
    monkeypatch.setattr(reports.returns,'fingerprint',Mock(side_effect=['before','after']))
    with pytest.raises(HTTPException) as error:
        reports.build(report_store,'2026-09','month','purchases',lambda _:profile)
    assert error.value.status_code==409


def test_samples_and_oversized_reports_fail_closed(report_store, invoice, profile, monkeypatch):
    invoice.sample=True
    with pytest.raises(HTTPException) as error:
        reports.build(report_store,'2026-09','month','purchases',lambda _:profile)
    assert error.value.status_code==403
    invoice.sample=False
    monkeypatch.setattr(reports,'MAX_ROWS',0)
    with pytest.raises(HTTPException) as error:
        reports.build(report_store,'2026-09','month','purchases',lambda _:profile)
    assert error.value.status_code==422


def test_audit_query_is_scoped_and_does_not_return_payloads(report_store):
    reports.audit_rows(report_store,'2026-09','2026-10-15T00:00:00+00:00')
    sql,params=report_store.query.call_args.args
    assert sql.count('client_id=@client_id')==3
    assert sql.count('created_at<=@cutoff')==3
    assert 'worker_token' not in sql and 'raw_ref' not in sql
    assert {item.name for item in params}=={'period','cutoff'}


def test_pdf_is_summary_only_and_invalid_formats_are_rejected(report_store, profile):
    result=reports.build(report_store,'2026-09','month','purchases',lambda _:profile)
    for format in ('pdf','html','exe'):
        with pytest.raises(HTTPException) as error:
            reports.render(result,format)
        assert error.value.status_code==422


def test_control_characters_are_preserved_visibly_and_long_cells_are_rejected():
    assert reports.safe_cell('A\x00B')=='A\\u0000B'
    with pytest.raises(HTTPException):
        reports.safe_cell('x'*32768)


def test_foreign_evidence_is_rejected(report_store, invoice, profile):
    invoice.client_id='other'
    with pytest.raises(HTTPException) as error:
        reports.build(report_store,'2026-09','month','invoices',lambda _:profile)
    assert error.value.status_code==403


def test_reports_keep_four_cash_scenarios_independent(report_store, invoice, profile):
    sale = invoice.model_copy(deep=True)
    sale.invoice_id = 'sale-control'
    sale.direction = 'sale'
    for head in ('taxable_value','cgst','sgst','invoice_total'):
        setattr(sale.totals, head, getattr(sale.totals, head)*2)
    for line in sale.line_items:
        for head in ('taxable_value','cgst','sgst'):
            setattr(line, head, getattr(line, head)*2)
    report_store.list.side_effect = None
    report_store.list.return_value = [invoice, sale]
    statement = {'supplier_gstin': invoice.supplier_gstin, 'invoice_number': invoice.invoice_number,
                 'invoice_date': str(invoice.invoice_date), 'itc_available': True,
                 **{head: str(getattr(invoice.totals, head)) for head in ('taxable_value','igst','cgst','sgst','cess')}}
    report_store.matching_rows.return_value = [statement]
    result = reports.build(report_store,'2026-09','month','tax',lambda _:profile)
    assert result['monthly_cash_scenarios'][0] == {
        'period':'2026-09','recorded':'180.00','matched':'180.00','eligible':'180.00','noCredit':'360.00'}
    report_store.matching_rows.return_value = []
    unmatched = reports.build(report_store,'2026-09','month','tax',lambda _:profile)
    assert unmatched['monthly_cash_scenarios'][0] == {
        'period':'2026-09','recorded':'180.00','matched':'360.00','eligible':'360.00','noCredit':'360.00'}


@pytest.mark.parametrize('kind', ['credit','liability','cash','without_itc','pending','mismatches','unmatched'])
def test_focused_reports_reuse_evidence_and_exports(report_store, profile, kind):
    result=reports.build(report_store,'2026-09','month',kind,lambda _:profile)
    assert result['client_id']=='a'
    assert result['kind']==kind
    for format in ('csv','xlsx'):
        content,mime=reports.render(result,format)
        assert content and mime


def test_focused_reports_match_dashboard_classifications(report_store, invoice, profile):
    from app.gst.report import calculate
    pending=invoice.model_copy(deep=True);pending.status='needs_review';pending.invoice_id='pending'
    rejected=invoice.model_copy(deep=True);rejected.status='rejected';rejected.invoice_id='rejected'
    report_store.list.side_effect=None
    report_store.list.return_value=[invoice,pending,rejected]
    result=reports.build(report_store,'2026-09','month','pending',lambda _:profile)
    assert [row[1] for row in result['rows']]==['pending']
    without=reports.build(report_store,'2026-09','month','without_itc',lambda _:profile)
    expected=calculate([invoice,pending,rejected],profile,[])
    from decimal import Decimal
    assert sum((Decimal(row[-1]) for row in without['rows'] if row[6]!='eligible'),Decimal('0'))==expected['totals']['non_gst']
    unmatched=reports.build(report_store,'2026-09','month','unmatched',lambda _:profile)
    assert all(row[6] in {'missing_book','missing_statement'} for row in unmatched['rows'])
    mismatches=reports.build(report_store,'2026-09','month','mismatches',lambda _:profile)
    assert all(row[6] not in {'exact','missing_book','missing_statement'} for row in mismatches['rows'])
    credit=reports.build(report_store,'2026-09','month','credit',lambda _:profile)
    assert all(row[6]=='eligible' for row in credit['rows'])
    assert credit['monthly_summary'][0]['totals']['eligible_credit']==format(expected['totals']['eligible_credit'],'.2f')


@pytest.mark.parametrize('kind', ['credit','liability','cash','without_itc','pending','mismatches','unmatched'])
def test_focused_reports_reject_foreign_evidence(report_store, invoice, profile, kind):
    invoice.client_id='foreign'
    report_store.list.side_effect=None;report_store.list.return_value=[invoice]
    with pytest.raises(HTTPException) as error:
        reports.build(report_store,'2026-09','month',kind,lambda _:profile)
    assert error.value.status_code==403


def test_focused_report_exports_retain_context_and_missing_data(report_store, profile):
    result=reports.build(report_store,'2026-08','month','credit',lambda _:profile)
    content,_=reports.render(result,'xlsx')
    book=load_workbook(BytesIO(content))
    assert book['Financial context']['D2'].value=='Not calculated'
    content,_=reports.render(result,'csv')
    assert b'Not calculated' in content
