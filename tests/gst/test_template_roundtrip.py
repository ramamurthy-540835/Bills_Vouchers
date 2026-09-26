from datetime import date
from decimal import Decimal
from io import BytesIO
from openpyxl import load_workbook
from app.gst.template import template_bytes,parse_template


def test_five_invoice_roundtrip(profile,invoice):
    workbook=load_workbook(BytesIO(template_bytes()))
    for i in range(5):
        for col,value in enumerate([invoice.supplier_gstin,f'BILL-{i}',date(2026,9,1),'8703','Car','1000.10','18','90.01','90.01','0','0','1180.12','33','N'],1):
            workbook['Bills'].cell(i+2,col,value)
    stream=BytesIO();workbook.save(stream)
    invoices=parse_template(stream.getvalue(),'.xlsx','a','d',profile,'gs://test/a/file')
    assert len(invoices)==5
    assert all(i.totals.taxable_value==Decimal('1000.10') and i.totals.invoice_total==Decimal('1180.12') for i in invoices)
    assert all(i.extraction.method=='xlsx_template' for i in invoices)


def test_csv_lines_group_by_invoice(profile,invoice):
    import csv
    from io import StringIO
    from app.gst.template import HEADERS
    stream=StringIO();writer=csv.writer(stream);writer.writerow(HEADERS)
    for _ in range(2):
        writer.writerow([invoice.supplier_gstin,'ONE','01-09-2026','8703','Car','100','18','9','9','0','0','236','33','N'])
    rows=parse_template(stream.getvalue().encode(),'.csv','a','d',profile,'ref')
    assert len(rows)==1 and len(rows[0].line_items)==2 and rows[0].totals.taxable_value==Decimal('200')
